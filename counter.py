from flask import Flask, jsonify, render_template, request, send_file, redirect
import requests
import json
from datetime import datetime, timedelta
import pytz
from collections import Counter, defaultdict
from apscheduler.schedulers.background import BackgroundScheduler
import os
import threading
import inspect
import re
import time
import sqlite3
import io
import csv
from openpyxl import Workbook

app = Flask(__name__)

def load_env():
    env_path = os.path.join(app.root_path, ".env")
    if not os.path.exists(env_path):
        return
    with open(env_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            value = value.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = value

load_env()

# === Конфигурация API ===
API_BASE         = "https://crm23.sipspeak.ru/api/shared"
CALL_LIST_URL    = f"{API_BASE}/call/list"
HIST_URL         = f"{API_BASE}/user_report/list/history"
LIST_URL         = f"{API_BASE}/user_report/list"
CONTACT_LIST_URL = f"{API_BASE}/contact/list"      # для новых номеров
CAMPAIGN_LIST_URL = f"{API_BASE}/campaign/list"    # для получения активных проектов
USERS_LIST_URL   = f"{API_BASE}/user/list"         # для получения операторов
API_TOKEN        = "cc7g45ybc7g5ync84umc9gmu5c9g4mucc"
HEADERS          = {"Authorization": API_TOKEN, "Accept": "application/json"}
REQUEST_TIMEOUT  = 60  # таймаут в секундах
PAGE_SIZE        = 500  # размер страницы

# === Google Sheets (ЦК) ===
CK_SHEET_CSV_URL = os.getenv("CK_SHEET_CSV_URL", "")

# === amoCRM (долгосрочный токен) ===
AMO_BASE_URL = os.getenv("AMO_BASE_URL", "").rstrip("/")
AMO_LONG_TOKEN = os.getenv("AMO_LONG_TOKEN", "")
AMO_HEADERS = {
    "Authorization": f"Bearer {AMO_LONG_TOKEN}",
    "Accept": "application/hal+json"
} if AMO_LONG_TOKEN else {}
AMO_USERS_CACHE_TTL = 3600
AMO_USERS_CACHE = {"ts": 0, "data": {}}
AMO_TZ = os.getenv("AMO_TZ", "Europe/Samara")
AMO_STATUS_MEETING_DONE = str(os.getenv("AMO_STATUS_MEETING_DONE", "79528826"))
AMO_STATUS_MEETING_OK = str(os.getenv("AMO_STATUS_MEETING_OK", "79652190"))
AMO_STATUS_DEAL_SUCCESS = str(os.getenv("AMO_STATUS_DEAL_SUCCESS", "142"))
AMO_FIELD_MEETING_OK = str(os.getenv("AMO_FIELD_MEETING_OK", "964369"))
AMO_FIELD_DEAL_SUM = str(os.getenv("AMO_FIELD_DEAL_SUM", "964601"))
AMO_FIELD_OPERATOR = str(os.getenv("AMO_FIELD_OPERATOR", "913949") or "913949")
AMO_FIELD_INV_CK = str(os.getenv("AMO_FIELD_INV_CK", "942511"))
AMO_FIELD_AGENT_CA = str(os.getenv("AMO_FIELD_AGENT_CA", "3635769") or "3635769")
AMO_INV_CK_VALUE = os.getenv("AMO_INV_CK_VALUE", "ЦК")
AMO_AGENT_CA_VALUE = os.getenv("AMO_AGENT_CA_VALUE", "ЦА")
AMO_INV_CK_ENUM_ID = str(os.getenv("AMO_INV_CK_ENUM_ID", "")).strip()
AMO_AGENT_CA_ENUM_ID = str(os.getenv("AMO_AGENT_CA_ENUM_ID", "")).strip()
AMO_PIPE_INVESTORS_ID = str(os.getenv("AMO_PIPE_INVESTORS_ID", "")).strip()
AMO_PIPE_AGENTS_ID = str(os.getenv("AMO_PIPE_AGENTS_ID", "9779410") or "9779410").strip()
AMO_PIPE_INVESTORS_NAME = os.getenv("AMO_PIPE_INVESTORS_NAME", "Воронка ФС")
AMO_PIPE_AGENTS_NAME = os.getenv("AMO_PIPE_AGENTS_NAME", "Запросы агентов")
AMO_DEBUG_EVENTS = os.getenv("AMO_DEBUG_EVENTS", "").lower() in ("1", "true", "yes", "y")

MOSCOW_OPERATOR_IDS = {oid.strip() for oid in (os.getenv("MOSCOW_OPERATOR_IDS", "38").split(",")) if oid.strip()}
MOSCOW_SIPSPEAK_AGREEMENTS_IDS = {oid.strip() for oid in (os.getenv("MOSCOW_SIPSPEAK_AGREEMENTS_IDS", "38").split(",")) if oid.strip()}
EXCLUDED_OPERATOR_IDS = {oid.strip() for oid in (os.getenv("EXCLUDED_OPERATOR_IDS", "38").split(",")) if oid.strip()}
AMO_CALL_MIN_SECONDS = int(os.getenv("AMO_CALL_MIN_SECONDS", "60"))
NIGHTLY_SYNC_TIME = os.getenv("CALLCENTER_NIGHTLY_SYNC_TIME", "02:30")
NIGHTLY_SYNC_DAYS = int(os.getenv("CALLCENTER_NIGHTLY_SYNC_DAYS", "0") or 0)
REPORT_AUTO_SYNC = os.getenv("REPORT_AUTO_SYNC", "0").strip().lower() in {"1", "true", "yes", "y"}
USE_ACTIVE_CAMPAIGNS = os.getenv("USE_ACTIVE_CAMPAIGNS", "0").strip().lower() in {"1", "true", "yes", "y"}

# === База данных ===
DB_PATH = os.getenv("CALLCENTER_DB_PATH", os.path.join(app.root_path, "data", "callcenter.db"))
ADMIN_TOKEN = os.getenv("CALLCENTER_ADMIN_TOKEN", "")


# === Операторы и статус-маппинг ===
OPERATORS = {}

STATUS_MAP = {
    "queue":      "В очереди",
    "active":     "Активный",
    "no_active":  "Неактивный",
    "pause":      "Перерыв",
    "call":       "Разговор",
    "card":       "Карточка",
    "dnd":        "Не беспокоить",
    "incoming":   "Входящий",
    "ringing":    "Звонит",
    "working_day":"Рабочий день"
}

ALL_CALLS_PARAMS = [("page",1), ("limit",10000)]
STAT_FULL = ["8","9","10","11","13","14","15","16","20","21","22","23","24","25","30","34","35","36","37","38","39"]
CS8        = ["8"]
CS20       = ["20"]
CS22       = ["22"]
LEAD_AGENT = ["22","30"]
AG_TRANSFER = ["36"]
AG_AGREEMENT = ["37", "22"]

def get_talk_duration(call):
    td = call.get("talk_duration")
    if td is None:
        td = call.get("duration")
    if td is None:
        td = call.get("billsec")
    if td is None:
        td = call.get("answered_seconds")
    if td is None:
        return 0
    if isinstance(td, (int, float)):
        try:
            return int(td)
        except (TypeError, ValueError):
            return 0

    raw = str(td).strip()
    if not raw:
        return 0

    # Форматы времени из SipSpeak встречаются как "MM:SS" и "HH:MM:SS".
    if ":" in raw:
        parts = raw.split(":")
        nums = []
        for p in parts:
            p = p.strip().replace(",", ".")
            try:
                nums.append(float(p))
            except (TypeError, ValueError):
                nums = []
                break
        if len(nums) == 2:
            return int(nums[0] * 60 + nums[1])
        if len(nums) == 3:
            return int(nums[0] * 3600 + nums[1] * 60 + nums[2])

    raw = raw.replace(",", ".")
    try:
        return int(float(raw))
    except (TypeError, ValueError):
        return 0

def get_status_id(call):
    status_obj = call.get("client_status")
    if isinstance(status_obj, dict):
        value = status_obj.get("id") or status_obj.get("value") or status_obj.get("status_id")
        if value is not None and value != "":
            return str(value)
    elif status_obj is not None and status_obj != "":
        return str(status_obj)

    for key in ("client_status_id", "status_id", "result_status_id"):
        value = call.get(key)
        if value is not None and value != "":
            return str(value)
    return ""

# для тестирования локально:
TEST_DATE = os.getenv("TEST_DATE")  # e.g. "14-05-2025"

# Глобальная переменная для хранения активных проектов
active_campaigns = []

def short_name(name):
    parts = [p for p in (name or "").split() if p]
    return " ".join(parts[:2]) if parts else name

def format_hms(seconds):
    try:
        total = int(seconds)
    except (TypeError, ValueError):
        total = 0
    h = total // 3600
    m = (total % 3600) // 60
    s = total % 60
    return f"{h:02d}:{m:02d}:{s:02d}"

def normalize_name(value):
    if not value:
        return ""
    return " ".join(value.replace("ё", "е").replace("Ё", "Е").lower().split())

def parse_names_list(value):
    raw = value or ""
    parts = re.split(r"[,\n;]+", raw)
    return [p.strip() for p in parts if p.strip()]

MOSCOW_OPERATOR_NAME_ALLOWLIST = {
    normalize_name(short_name(name))
    for name in parse_names_list(
        os.getenv(
            "MOSCOW_OPERATOR_NAMES",
            "Самаркин Александр,Чисникова Валерия,Шарыхин Александр,Максим Клочков"
        )
    )
}

def moscow_name_allowed(name):
    if not MOSCOW_OPERATOR_NAME_ALLOWLIST:
        return True
    return normalize_name(short_name(name)) in MOSCOW_OPERATOR_NAME_ALLOWLIST

def parse_sheet_date(value):
    raw = (value or "").strip()
    if not raw:
        return None
    for fmt in ("%d.%m.%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None

@app.route("/amo/lead/<int:lead_id>")
def amo_lead_debug(lead_id):
    if ADMIN_TOKEN and request.args.get("token") != ADMIN_TOKEN:
        return jsonify({"ok": False, "error": "unauthorized"}), 401
    if not amo_enabled():
        return jsonify({"ok": False, "error": "amo disabled"}), 400
    try:
        r = amo_get(f"/api/v4/leads/{lead_id}")
    except requests.RequestException as e:
        print(f"AMO lead {lead_id} request error: {e}")
        return jsonify({"ok": False, "error": str(e)}), 502
    print(f"AMO lead {lead_id} status: {r.status_code}")
    content_type = r.headers.get("Content-Type", "")
    print(f"AMO lead {lead_id} content-type: {content_type}")
    if r.status_code >= 400:
        print(f"AMO lead {lead_id} body: {r.text}")
        return jsonify({"ok": False, "status": r.status_code, "body": r.text}), 502
    try:
        data = r.json()
    except ValueError:
        print(f"AMO lead {lead_id} invalid JSON: {r.text}")
        return jsonify({"ok": False, "error": "invalid json"}), 502
    print(f"AMO lead {lead_id} response:\n{data}")
    return jsonify(data)

def amo_enabled():
    return bool(AMO_BASE_URL and AMO_LONG_TOKEN)

def amo_get(path, params=None):
    url = f"{AMO_BASE_URL}{path}"
    return requests.get(url, headers=AMO_HEADERS, params=params, timeout=REQUEST_TIMEOUT)

def amo_users_fallback():
    path = os.path.join(app.root_path, "users_map.json")
    if not os.path.exists(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return {str(k): v for k, v in data.items()}
    except Exception as e:
        print(f"AMO users fallback read error: {e}")
        return {}

def amo_fetch_users():
    if not amo_enabled():
        return {}
    now = time.time()
    if AMO_USERS_CACHE["data"] and now - AMO_USERS_CACHE["ts"] < AMO_USERS_CACHE_TTL:
        return AMO_USERS_CACHE["data"]
    users = {}
    page = 1
    debug_seen = 0
    while True:
        try:
            r = amo_get("/api/v4/users", params={"limit": 250, "page": page})
        except requests.RequestException as e:
            print(f"AMO users request error: {e}")
            break
        if r.status_code == 403:
            print("AMO users 403 - using users_map.json fallback")
            users = amo_users_fallback()
            break
        if r.status_code != 200:
            print(f"AMO users HTTP {r.status_code}: {r.text}")
            break
        data = r.json()
        batch = data.get("_embedded", {}).get("users", []) or []
        if not batch:
            break
        for user in batch:
            uid = str(user.get("id") or "")
            if uid:
                users[uid] = user.get("name") or uid
        page += 1
    AMO_USERS_CACHE["ts"] = now
    AMO_USERS_CACHE["data"] = users
    return users

AMO_PIPELINES_CACHE = {"ts": 0, "name_to_id": {}}

def amo_fetch_pipelines():
    if not amo_enabled():
        return {}
    now = time.time()
    if AMO_PIPELINES_CACHE["name_to_id"] and now - AMO_PIPELINES_CACHE["ts"] < AMO_USERS_CACHE_TTL:
        return AMO_PIPELINES_CACHE["name_to_id"]
    out = {}
    page = 1
    seen_pages = set()
    while True:
        if page in seen_pages:
            break
        seen_pages.add(page)
        try:
            r = amo_get("/api/v4/leads/pipelines", params={"limit": 250, "page": page})
        except requests.RequestException as e:
            print(f"AMO pipelines request error: {e}")
            break
        if r.status_code == 204:
            break
        if r.status_code != 200:
            print(f"AMO pipelines HTTP {r.status_code}: {r.text}")
            break
        data = r.json()
        pipes = data.get("_embedded", {}).get("pipelines", []) or []
        if not pipes:
            break
        for pipe in pipes:
            pid = str(pipe.get("id") or "").strip()
            name = pipe.get("name") or ""
            if pid and name:
                out[normalize_name(name)] = pid
        links = data.get("_links") or {}
        if not links.get("next"):
            break
        page += 1
    AMO_PIPELINES_CACHE["ts"] = now
    AMO_PIPELINES_CACHE["name_to_id"] = out
    return out

def amo_resolve_pipeline_id(pipeline_id, pipeline_name):
    if pipeline_id:
        return str(pipeline_id).strip()
    lookup = amo_fetch_pipelines()
    return lookup.get(normalize_name(pipeline_name), "")

def amo_day_range(date_str):
    tz = pytz.timezone(AMO_TZ)
    d = parse_date(date_str)
    start = tz.localize(datetime(d.year, d.month, d.day, 0, 0, 0))
    end = tz.localize(datetime(d.year, d.month, d.day, 23, 59, 59))
    return int(start.timestamp()), int(end.timestamp())

def amo_range_timestamps(start_date, end_date):
    if not start_date or not end_date:
        return None, None
    if parse_date(start_date) > parse_date(end_date):
        start_date, end_date = end_date, start_date
    tz = pytz.timezone(AMO_TZ)
    d_start = parse_date(start_date)
    d_end = parse_date(end_date)
    start = tz.localize(datetime(d_start.year, d_start.month, d_start.day, 0, 0, 0))
    end = tz.localize(datetime(d_end.year, d_end.month, d_end.day, 23, 59, 59))
    return int(start.timestamp()), int(end.timestamp())

def amo_find_custom_field(lead, field_id):
    for field in (lead.get("custom_fields_values") or []):
        if str(field.get("field_id") or "") == str(field_id):
            return field
    return None

def amo_field_checkbox_true(lead, field_id):
    field = amo_find_custom_field(lead, field_id)
    if not field:
        return False
    for item in (field.get("values") or []):
        value = item.get("value")
        if value is True:
            return True
        if isinstance(value, str) and value.strip().lower() in {"true", "1", "yes", "да"}:
            return True
        if isinstance(value, (int, float)) and value == 1:
            return True
    return False

def amo_field_numeric(lead, field_id):
    field = amo_find_custom_field(lead, field_id)
    if not field:
        return 0
    value = None
    for item in (field.get("values") or []):
        value = item.get("value")
        if value is not None:
            break
    if value is None:
        return 0
    try:
        return int(float(str(value).replace(" ", "").replace(",", ".")))
    except (TypeError, ValueError):
        return 0

def amo_field_text(lead, field_id):
    field = amo_find_custom_field(lead, field_id)
    if not field:
        return ""
    for item in (field.get("values") or []):
        value = item.get("value")
        if value is None:
            value = item.get("text")
        if value is None:
            continue
        text = str(value).strip()
        if text:
            return text
    return ""

def amo_tokens_match(tokens, value_text="", enum_id=""):
    if not tokens:
        return False
    for token in tokens:
        if enum_id and token.get("enum_id") == enum_id:
            return True
        txt = normalize_name(token.get("text", ""))
        if value_text and txt == value_text:
            return True
    return False

def amo_events_call_notes(date_str):
    if not amo_enabled():
        return {"leads": [], "contacts": []}
    start_ts, end_ts = amo_day_range(date_str)
    lead_note_ids = []
    contact_note_ids = []
    page = 1
    total = 0
    while True:
        params = [
            ("limit", 250),
            ("page", page),
            ("filter[type][]", "incoming_call"),
            ("filter[type][]", "outgoing_call"),
            ("filter[created_at][from]", start_ts),
            ("filter[created_at][to]", end_ts),
        ]
        try:
            r = amo_get("/api/v4/events", params=params)
        except requests.RequestException as e:
            print(f"AMO events request error: {e}")
            break
        if r.status_code == 204:
            break
        if r.status_code == 204:
            break
        if r.status_code != 200:
            print(f"AMO events HTTP {r.status_code}: {r.text}")
            break
        data = r.json()
        events = data.get("_embedded", {}).get("events", []) or []
        if not events:
            break
        for event in events:
            total += 1
            note_id = None
            for item in (event.get("value_after") or []):
                note = item.get("note") or {}
                note_id = note.get("id")
                if note_id:
                    break
            if not note_id:
                continue
            entity_type = event.get("entity_type")
            if entity_type == "lead":
                lead_note_ids.append(note_id)
            elif entity_type == "contact":
                contact_note_ids.append(note_id)
        page += 1
    print(f"AMO call events {date_str}: {total}")
    return {
        "leads": list(dict.fromkeys(lead_note_ids)),
        "contacts": list(dict.fromkeys(contact_note_ids))
    }

def amo_fetch_notes(entity_type, note_ids):
    if not note_ids:
        return []
    endpoint = f"/api/v4/{entity_type}/notes"
    notes = []
    chunk_size = 200
    for i in range(0, len(note_ids), chunk_size):
        chunk = note_ids[i:i+chunk_size]
        params = [("limit", 250)]
        for note_id in chunk:
            params.append(("filter[id][]", note_id))
        try:
            r = amo_get(endpoint, params=params)
        except requests.RequestException as e:
            print(f"AMO notes request error: {e}")
            continue
        if r.status_code != 200:
            print(f"AMO notes HTTP {r.status_code}: {r.text}")
            continue
        data = r.json()
        batch = data.get("_embedded", {}).get("notes", []) or []
        notes.extend(batch)
    return notes

def amo_calls_over_minute(date_str):
    if not amo_enabled():
        return Counter()
    note_ids = amo_events_call_notes(date_str)
    lead_notes = amo_fetch_notes("leads", note_ids["leads"])
    contact_notes = amo_fetch_notes("contacts", note_ids["contacts"])
    counts = Counter()

    def handle(note):
        note_type = note.get("note_type") or ""
        if note_type not in ("call_in", "call_out"):
            return
        params = note.get("params") or {}
        try:
            duration = int(params.get("duration") or 0)
        except (TypeError, ValueError):
            duration = 0
        if duration < AMO_CALL_MIN_SECONDS:
            return
        rid = str(note.get("responsible_user_id") or "")
        if rid:
            counts[rid] += 1

    for note in lead_notes:
        handle(note)
    for note in contact_notes:
        handle(note)
    print(f"AMO calls >= {AMO_CALL_MIN_SECONDS}s {date_str}: {sum(counts.values())}")
    return counts

def amo_events_call_notes_range(start_date, end_date):
    if not amo_enabled():
        return {"leads": [], "contacts": []}
    start_ts, end_ts = amo_range_timestamps(start_date, end_date)
    if not start_ts or not end_ts:
        return {"leads": [], "contacts": []}
    lead_note_ids = []
    contact_note_ids = []
    page = 1
    total = 0
    while True:
        params = [
            ("limit", 250),
            ("page", page),
            ("filter[type][]", "incoming_call"),
            ("filter[type][]", "outgoing_call"),
            ("filter[created_at][from]", start_ts),
            ("filter[created_at][to]", end_ts),
        ]
        try:
            r = amo_get("/api/v4/events", params=params)
        except requests.RequestException as e:
            print(f"AMO events request error: {e}")
            break
        if r.status_code != 200:
            print(f"AMO events HTTP {r.status_code}: {r.text}")
            break
        data = r.json()
        events = data.get("_embedded", {}).get("events", []) or []
        if not events:
            break
        for event in events:
            total += 1
            note_id = None
            for item in (event.get("value_after") or []):
                note = item.get("note") or {}
                note_id = note.get("id")
                if note_id:
                    break
            if not note_id:
                continue
            entity_type = event.get("entity_type")
            if entity_type == "lead":
                lead_note_ids.append(note_id)
            elif entity_type == "contact":
                contact_note_ids.append(note_id)
        page += 1
    print(f"AMO call events {start_date}..{end_date}: {total}")
    return {
        "leads": list(dict.fromkeys(lead_note_ids)),
        "contacts": list(dict.fromkeys(contact_note_ids))
    }

def amo_calls_over_minute_range(start_date, end_date):
    if not amo_enabled():
        return Counter()
    note_ids = amo_events_call_notes_range(start_date, end_date)
    lead_notes = amo_fetch_notes("leads", note_ids["leads"])
    contact_notes = amo_fetch_notes("contacts", note_ids["contacts"])
    counts = Counter()

    def handle(note):
        note_type = note.get("note_type") or ""
        if note_type not in ("call_in", "call_out"):
            return
        params = note.get("params") or {}
        try:
            duration = int(params.get("duration") or 0)
        except (TypeError, ValueError):
            duration = 0
        if duration < AMO_CALL_MIN_SECONDS:
            return
        rid = str(note.get("responsible_user_id") or "")
        if rid:
            counts[rid] += 1

    for note in lead_notes:
        handle(note)
    for note in contact_notes:
        handle(note)
    print(f"AMO calls >= {AMO_CALL_MIN_SECONDS}s {start_date}..{end_date}: {sum(counts.values())}")
    return counts

def amo_leads_created_metrics(date_str):
    if not amo_enabled():
        return {
            "agreement": Counter(),
            "meeting": Counter(),
            "success": Counter(),
            "revenue": defaultdict(int)
        }
    start_ts, end_ts = amo_day_range(date_str)
    agreement = Counter()
    meeting = Counter()
    success = Counter()
    revenue = defaultdict(int)
    page = 1
    total = 0
    while True:
        params = {
            "limit": 250,
            "page": page,
            "filter[created_at][from]": start_ts,
            "filter[created_at][to]": end_ts
        }
        try:
            r = amo_get("/api/v4/leads", params=params)
        except requests.RequestException as e:
            print(f"AMO leads request error: {e}")
            break
        if r.status_code != 200:
            print(f"AMO leads HTTP {r.status_code}: {r.text}")
            break
        data = r.json()
        leads = data.get("_embedded", {}).get("leads", []) or []
        if not leads:
            break
        for lead in leads:
            rid = str(lead.get("responsible_user_id") or "")
            if rid:
                total += 1
                status_id = str(lead.get("status_id") or "")
                if amo_field_checkbox_true(lead, AMO_FIELD_MEETING_OK):
                    agreement[rid] += 1
                if status_id == AMO_STATUS_MEETING_DONE:
                    meeting[rid] += 1
                if status_id == AMO_STATUS_DEAL_SUCCESS:
                    success[rid] += 1
                amount = amo_field_numeric(lead, AMO_FIELD_DEAL_SUM)
                if amount:
                    revenue[rid] += amount
        page += 1
    print(f"AMO leads {date_str}: {total}")
    return {
        "agreement": agreement,
        "meeting": meeting,
        "success": success,
        "revenue": revenue
    }

def amo_leads_event_metrics(date_str):
    if not amo_enabled():
        return {
            "agreement": Counter(),
            "meeting": Counter(),
            "success": Counter(),
            "revenue": defaultdict(int)
        }
    start_ts, end_ts = amo_day_range(date_str)
    agreement_latest = {}
    status_latest = {}
    created_by_map = {}
    page = 1
    total = 0
    debug_seen = 0

    def extract_field_change(event):
        value_after = event.get("value_after")
        value_before = event.get("value_before")

        if isinstance(value_after, list) and len(value_after) == 0 and value_before:
            items = value_before if isinstance(value_before, list) else [value_before]
            for item in items:
                if not isinstance(item, dict):
                    continue
                if isinstance(item.get("custom_field_value"), dict):
                    cf = item["custom_field_value"]
                    field_id = cf.get("field_id") or cf.get("id")
                    if field_id is not None:
                        return str(field_id), []
                field_id = item.get("field_id") or item.get("id") or item.get("custom_field_id")
                if field_id is not None:
                    return str(field_id), []
        def extract_from_items(items, fallback_empty=False):
            for item in items:
                if not isinstance(item, dict):
                    continue
                if isinstance(item.get("custom_field_value"), dict):
                    cf = item["custom_field_value"]
                    field_id = cf.get("field_id") or cf.get("id")
                    value = cf.get("value")
                    if value is None:
                        value = cf.get("text")
                    if value is None:
                        value = cf.get("enum_id")
                    values = [value] if value is not None else ([] if fallback_empty else None)
                    return (str(field_id) if field_id is not None else ""), values
                field_id = item.get("field_id") or item.get("id") or item.get("custom_field_id")
                if field_id is None:
                    field = item.get("field") or item.get("custom_field")
                    if isinstance(field, dict):
                        field_id = field.get("id") or field.get("field_id")
                if field_id is None:
                    field_id = event.get("field_id") or event.get("custom_field_id")
                values = item.get("values")
                if values is None and isinstance(item.get("value"), list):
                    values = item.get("value")
                if values is None and "value" in item:
                    values = [item.get("value")]
                return (str(field_id) if field_id is not None else ""), values
            return "", None

        items = []
        if isinstance(value_after, list):
            items = value_after
        elif isinstance(value_after, dict):
            items = [value_after]
        field_id, values = extract_from_items(items)
        if field_id:
            return field_id, values

        value_before = event.get("value_before")
        items = []
        if isinstance(value_before, list):
            items = value_before
        elif isinstance(value_before, dict):
            items = [value_before]
        field_id, values = extract_from_items(items, fallback_empty=True)
        if field_id:
            return field_id, values

        field_id = event.get("field_id") or event.get("custom_field_id")
        if field_id is not None:
            return str(field_id), event.get("values") or event.get("value")
        return "", None

    def values_true(values):
        if values is None:
            return False
        if not isinstance(values, list):
            values = [values]
        for v in values:
            if isinstance(v, dict):
                v = v.get("value")
            if isinstance(v, str):
                vv = v.strip().lower()
                if vv in ("true", "1", "да", "yes", "y", "on", "вкл", "вкл."):
                    return True
                if vv in ("false", "0", "нет", "no", "n", "off", "выкл", "выкл."):
                    return False
            if v is True or v == 1:
                return True
        return False

    def is_lead_event(event):
        entity_type = event.get("entity_type")
        if not entity_type:
            entity_type = (event.get("_embedded", {}) or {}).get("entity", {}) or {}
            entity_type = entity_type.get("type")
        if isinstance(entity_type, str):
            entity_type = entity_type.lower()
        if entity_type and entity_type not in ("lead", "leads"):
            return False
        return True

    def fetch_events(types, handler):
        page = 1
        while True:
            params = [
                ("limit", 250),
                ("page", page),
                ("filter[created_at][from]", start_ts),
                ("filter[created_at][to]", end_ts),
            ]
            for t in types:
                params.append(("filter[type][]", t))
            try:
                r = amo_get("/api/v4/events", params=params)
            except requests.RequestException as e:
                print(f"AMO events request error: {e}")
                break
            if r.status_code == 204:
                break
            if r.status_code != 200:
                print(f"AMO events HTTP {r.status_code}: {r.text}")
                break
            data = r.json()
            events = data.get("_embedded", {}).get("events", []) or []
            if not events:
                break
            for event in events:
                handler(event)
            page += 1

    def handle_status_event(event):
        nonlocal total
        if not is_lead_event(event):
            return
        lead_id = event.get("entity_id")
        if not lead_id:
            return
        total += 1
        event_ts = event.get("created_at") or 0
        created_by = str(event.get("created_by") or "")
        if created_by:
            created_by_map[lead_id] = created_by
        status_id = ""
        value_after = event.get("value_after")
        if isinstance(value_after, dict):
            status_id = str(value_after.get("status_id") or "")
            if not status_id and isinstance(value_after.get("lead_status"), dict):
                status_id = str(value_after["lead_status"].get("id") or "")
        elif isinstance(value_after, list):
            for item in value_after:
                if not isinstance(item, dict):
                    continue
                if "status_id" in item:
                    status_id = str(item.get("status_id") or "")
                    break
                if isinstance(item.get("lead_status"), dict):
                    status_id = str(item["lead_status"].get("id") or "")
                    if status_id:
                        break
        if status_id:
            prev = status_latest.get(lead_id)
            if not prev or event_ts >= prev[0]:
                status_latest[lead_id] = (event_ts, status_id)
            # keep only the latest status per lead; we'll decide counts from it

    def handle_field_event(event):
        nonlocal total, debug_seen
        if not is_lead_event(event):
            return
        lead_id = event.get("entity_id")
        if not lead_id:
            return
        total += 1
        event_ts = event.get("created_at") or 0
        created_by = str(event.get("created_by") or "")
        if created_by:
            created_by_map[lead_id] = created_by
        field_id, values = extract_field_change(event)
        if AMO_DEBUG_EVENTS and debug_seen < 20:
            print(f"AMO field change lead={lead_id} field_id={field_id} values={values}")
            debug_seen += 1
        if field_id == AMO_FIELD_MEETING_OK:
            flag = values_true(values)
            if AMO_DEBUG_EVENTS:
                print(f"AMO meeting_ok lead={lead_id} at={event_ts} values={values} flag={flag}")
            prev = agreement_latest.get(lead_id)
            if not prev or event_ts >= prev[0]:
                agreement_latest[lead_id] = (event_ts, flag)

    fetch_events(["lead_status_changed"], handle_status_event)
    fetch_events(["custom_field_value_changed", f"custom_field_{AMO_FIELD_MEETING_OK}_value_changed"], handle_field_event)

    lead_ids = set(agreement_latest.keys()) | set(status_latest.keys())
    lead_ids = list(lead_ids)

    lead_map = {}
    lead_amounts = {}
    chunk_size = 250
    for i in range(0, len(lead_ids), chunk_size):
        chunk = lead_ids[i:i+chunk_size]
        params = [("limit", 250)]
        for lid in chunk:
            params.append(("filter[id][]", str(lid)))
        try:
            r = amo_get("/api/v4/leads", params=params)
        except requests.RequestException as e:
            print(f"AMO leads request error: {e}")
            continue
        if r.status_code == 204:
            continue
        if r.status_code != 200:
            print(f"AMO leads HTTP {r.status_code}: {r.text}")
            continue
        data = r.json()
        leads = data.get("_embedded", {}).get("leads", []) or []
        for lead in leads:
            lead_id = lead.get("id")
            rid = str(lead.get("responsible_user_id") or "")
            if lead_id and rid:
                lead_map[lead_id] = rid
                lead_amounts[lead_id] = amo_field_numeric(lead, AMO_FIELD_DEAL_SUM)

    agreement = Counter()
    meeting = Counter()
    success = Counter()
    revenue = defaultdict(int)

    for lid, (ts, flag) in agreement_latest.items():
        if not flag:
            continue
        rid = lead_map.get(lid) or created_by_map.get(lid)
        if rid:
            agreement[rid] += 1
    # Count unique leads that reached the status during the day
    for lid, (ts, status_id) in status_latest.items():
        rid = lead_map.get(lid) or created_by_map.get(lid)
        if not rid:
            continue
        if status_id == AMO_STATUS_MEETING_DONE:
            meeting[rid] += 1
        if status_id == AMO_STATUS_DEAL_SUCCESS:
            success[rid] += 1
            amount = lead_amounts.get(lid, 0)
            if amount:
                revenue[rid] += amount

    print(f"AMO lead events {date_str}: {total}")
    return {
        "agreement": agreement,
        "meeting": meeting,
        "success": success,
        "revenue": revenue
    }

def amo_ulya_invest_agent_metrics(date_str, operators_override=None):
    if not amo_enabled():
        return {"inv_ck": Counter(), "ag_ca": Counter()}

    inv_pipeline_id = amo_resolve_pipeline_id(AMO_PIPE_INVESTORS_ID, AMO_PIPE_INVESTORS_NAME)
    ag_pipeline_id = amo_resolve_pipeline_id(AMO_PIPE_AGENTS_ID, AMO_PIPE_AGENTS_NAME)
    field_specs = []
    if AMO_FIELD_AGENT_CA:
        field_specs.append({
            "key": "ag_ca",
            "field_id": str(AMO_FIELD_AGENT_CA),
            "value_text": normalize_name(AMO_AGENT_CA_VALUE),
            "enum_id": AMO_AGENT_CA_ENUM_ID,
            "pipeline_id": ag_pipeline_id,
            "use_operator_field": True
        })
    if not field_specs:
        return {"inv_ck": Counter(), "ag_ca": Counter()}

    start_ts, end_ts = amo_day_range(date_str)
    latest = {spec["key"]: {} for spec in field_specs}
    created_by_map = {}

    def extract_change(event):
        value_after = event.get("value_after")
        value_before = event.get("value_before")

        if isinstance(value_after, list) and len(value_after) == 0:
            items = value_before if isinstance(value_before, list) else [value_before]
            for item in items:
                if not isinstance(item, dict):
                    continue
                cf = item.get("custom_field_value")
                if isinstance(cf, dict):
                    fid = cf.get("field_id") or cf.get("id")
                    if fid is not None:
                        return str(fid), []
                fid = item.get("field_id") or item.get("id") or item.get("custom_field_id")
                if fid is not None:
                    return str(fid), []

        items = value_after if isinstance(value_after, list) else ([value_after] if isinstance(value_after, dict) else [])
        for item in items:
            if not isinstance(item, dict):
                continue
            cf = item.get("custom_field_value")
            if isinstance(cf, dict):
                fid = cf.get("field_id") or cf.get("id")
                if fid is None:
                    continue
                token = {
                    "text": str(cf.get("text") if cf.get("text") is not None else (cf.get("value") if cf.get("value") is not None else "")).strip(),
                    "enum_id": str(cf.get("enum_id") or "").strip()
                }
                return str(fid), [token]
            fid = item.get("field_id") or item.get("id") or item.get("custom_field_id")
            if fid is None:
                continue
            vals = item.get("values")
            tokens = []
            if isinstance(vals, list):
                for v in vals:
                    if isinstance(v, dict):
                        tokens.append({
                            "text": str(v.get("value") if v.get("value") is not None else "").strip(),
                            "enum_id": str(v.get("enum_id") or "").strip()
                        })
                    else:
                        tokens.append({"text": str(v).strip(), "enum_id": ""})
            elif "value" in item:
                tokens.append({"text": str(item.get("value") if item.get("value") is not None else "").strip(), "enum_id": ""})
            return str(fid), tokens

        return "", []

    def is_lead_event(event):
        entity_type = (event.get("entity_type") or "").lower()
        return entity_type in ("lead", "leads")

    def fetch_events_for(spec):
        page = 1
        while True:
            params = [
                ("limit", 250),
                ("page", page),
                ("filter[created_at][from]", start_ts),
                ("filter[created_at][to]", end_ts),
                ("filter[type][]", "custom_field_value_changed"),
            ]
            try:
                r = amo_get("/api/v4/events", params=params)
            except requests.RequestException as e:
                print(f"AMO events request error: {e}")
                break
            if r.status_code == 204:
                break
            if r.status_code == 400:
                # Некоторые аккаунты amoCRM не поддерживают типы вида custom_field_<id>_value_changed.
                # Используем только универсальный custom_field_value_changed и фильтруем по field_id в коде.
                print(f"AMO events HTTP 400 for field {spec['field_id']} ({spec['key']}): {r.text}")
                break
            if r.status_code != 200:
                print(f"AMO events HTTP {r.status_code}: {r.text}")
                break
            data = r.json()
            events = data.get("_embedded", {}).get("events", []) or []
            if not events:
                break
            for event in events:
                if not is_lead_event(event):
                    continue
                lead_id = event.get("entity_id")
                if not lead_id:
                    continue
                created_by = str(event.get("created_by") or "")
                if created_by:
                    created_by_map[lead_id] = created_by
                fid, tokens = extract_change(event)
                if fid != spec["field_id"]:
                    continue
                ts = event.get("created_at") or 0
                prev = latest[spec["key"]].get(lead_id)
                if not prev or ts >= prev[0]:
                    latest[spec["key"]][lead_id] = (ts, tokens)
            page += 1

    for spec in field_specs:
        fetch_events_for(spec)

    lead_ids = set()
    for spec in field_specs:
        lead_ids |= set(latest[spec["key"]].keys())
    lead_ids = list(lead_ids)

    source_map = operators_override or OPERATORS or {}
    operators_source = {
        str(oid): name
        for oid, name in source_map.items()
        if str(oid) not in EXCLUDED_OPERATOR_IDS and str(oid) not in MOSCOW_OPERATOR_IDS
    }
    if not operators_source:
        operators_source = {
            str(oid): name for oid, name in (OPERATORS or {}).items()
            if str(oid) not in EXCLUDED_OPERATOR_IDS
        }
    operator_name_map = build_operator_name_map(operators_source)
    unknown_operator_names = set()

    lead_meta = {}
    for i in range(0, len(lead_ids), 250):
        chunk = lead_ids[i:i+250]
        params = [("limit", 250)]
        for lid in chunk:
            params.append(("filter[id][]", str(lid)))
        try:
            r = amo_get("/api/v4/leads", params=params)
        except requests.RequestException as e:
            print(f"AMO leads request error: {e}")
            continue
        if r.status_code == 204:
            continue
        if r.status_code != 200:
            print(f"AMO leads HTTP {r.status_code}: {r.text}")
            continue
        data = r.json()
        leads = data.get("_embedded", {}).get("leads", []) or []
        for lead in leads:
            lid = lead.get("id")
            if not lid:
                continue
            operator_name = amo_field_text(lead, AMO_FIELD_OPERATOR)
            operator_id = ""
            if operator_name:
                key_full = normalize_name(operator_name)
                key_short = normalize_name(short_name(operator_name))
                operator_id = operator_name_map.get(key_full) or operator_name_map.get(key_short) or ""
                if not operator_id:
                    unknown_operator_names.add(operator_name)
            lead_meta[lid] = {
                "responsible": str(lead.get("responsible_user_id") or ""),
                "pipeline_id": str(lead.get("pipeline_id") or ""),
                "operator_id": str(operator_id or "")
            }

    result = {"inv_ck": Counter(), "ag_ca": Counter()}
    for spec in field_specs:
        key = spec["key"]
        target_pipeline = spec["pipeline_id"]
        for lid, (_, event_tokens) in latest[key].items():
            if not amo_tokens_match(
                event_tokens,
                value_text=spec.get("value_text", ""),
                enum_id=spec.get("enum_id", "")
            ):
                continue
            meta = lead_meta.get(lid, {})
            if target_pipeline and str(meta.get("pipeline_id") or "") != str(target_pipeline):
                continue
            if spec.get("use_operator_field"):
                rid = str(meta.get("operator_id") or "")
                if not rid:
                    rid = str(meta.get("responsible") or "") or created_by_map.get(lid, "")
            else:
                rid = str(meta.get("responsible") or "") or created_by_map.get(lid, "")
            if rid:
                result[key][rid] += 1

    if unknown_operator_names:
        print(f"AMO ЦА: operator name not mapped: {', '.join(sorted(unknown_operator_names))}")
    print(f"AMO ULYA {date_str}: ag_ca={dict(result['ag_ca'])}")

    return result

def amo_leads_created_metrics_range(start_date, end_date):
    if not amo_enabled():
        return {
            "agreement": Counter(),
            "meeting": Counter(),
            "success": Counter(),
            "revenue": defaultdict(int)
        }
    start_ts, end_ts = amo_range_timestamps(start_date, end_date)
    if not start_ts or not end_ts:
        return {
            "agreement": Counter(),
            "meeting": Counter(),
            "success": Counter(),
            "revenue": defaultdict(int)
        }
    agreement = Counter()
    meeting = Counter()
    success = Counter()
    revenue = defaultdict(int)
    page = 1
    total = 0
    while True:
        params = {
            "limit": 250,
            "page": page,
            "filter[created_at][from]": start_ts,
            "filter[created_at][to]": end_ts
        }
        try:
            r = amo_get("/api/v4/leads", params=params)
        except requests.RequestException as e:
            print(f"AMO leads request error: {e}")
            break
        if r.status_code != 200:
            print(f"AMO leads HTTP {r.status_code}: {r.text}")
            break
        data = r.json()
        leads = data.get("_embedded", {}).get("leads", []) or []
        if not leads:
            break
        for lead in leads:
            total += 1
            if lead.get("is_deleted"):
                continue
            rid = str(lead.get("responsible_user_id") or "")
            if not rid:
                continue
            status_id = str(lead.get("status_id") or "")
            if amo_field_checkbox_true(lead, AMO_FIELD_MEETING_OK):
                agreement[rid] += 1
            if status_id == AMO_STATUS_MEETING_DONE:
                meeting[rid] += 1
            if status_id == AMO_STATUS_DEAL_SUCCESS:
                success[rid] += 1
            amount = amo_field_numeric(lead, AMO_FIELD_DEAL_SUM)
            if amount:
                revenue[rid] += amount
        page += 1
    print(f"AMO leads {start_date}..{end_date}: {total}")
    return {
        "agreement": agreement,
        "meeting": meeting,
        "success": success,
        "revenue": revenue
    }

def merge_amo_counts(payload, date_str):
    if not amo_enabled():
        return payload
    amo_metrics = amo_leads_event_metrics(date_str)
    ulya_amo = amo_ulya_invest_agent_metrics(date_str, operators_override=payload.get("operators") or OPERATORS)
    amo_calls_1m = amo_calls_over_minute(date_str)
    if not amo_metrics or not (amo_metrics["success"] or amo_metrics["agreement"] or amo_metrics["meeting"] or amo_metrics["revenue"]):
        payload["amo_deals"] = {}
        payload["amo_operators"] = {}
        payload["amo_agreements"] = {}
        payload["amo_meetings"] = {}
        payload["amo_success"] = {}
        payload["amo_revenue"] = {}
        payload["amo_calls_1m"] = {oid: amo_calls_1m.get(oid, 0) for oid in amo_calls_1m}
        payload["inv_ck"] = {oid: (payload.get("ck_lead", {}) or {}).get(oid, 0) for oid in (payload.get("operators") or {})}
        payload["ag_ca"] = {oid: ulya_amo["ag_ca"].get(oid, 0) for oid in (payload.get("operators") or {})}
        return payload
    amo_users = amo_fetch_users()
    all_ids = set(amo_metrics["success"]) | set(amo_metrics["agreement"]) | set(amo_metrics["meeting"]) | set(amo_metrics["revenue"])
    all_ids |= set(amo_calls_1m)
    amo_ops = {oid: short_name(amo_users.get(oid, f"User {oid}")) for oid in all_ids if oid not in EXCLUDED_OPERATOR_IDS}
    if not MOSCOW_OPERATOR_NAME_ALLOWLIST:
        amo_name_set = set(amo_ops.values())
        for oid in MOSCOW_OPERATOR_IDS:
            if oid in EXCLUDED_OPERATOR_IDS:
                continue
            name = short_name(OPERATORS.get(oid) or f"User {oid}")
            if oid not in amo_ops and name not in amo_name_set:
                amo_ops[oid] = name
    if MOSCOW_OPERATOR_NAME_ALLOWLIST:
        amo_ops = {oid: name for oid, name in amo_ops.items() if moscow_name_allowed(name)}
    operators = payload.get("operators") or {}
    all_ids = set(amo_ops) if MOSCOW_OPERATOR_NAME_ALLOWLIST else (set(operators) | set(amo_ops))
    payload["amo_operators"] = amo_ops
    payload["amo_deals"] = {oid: amo_metrics["success"].get(oid, 0) for oid in all_ids if oid not in EXCLUDED_OPERATOR_IDS}
    payload["amo_agreements"] = {oid: amo_metrics["agreement"].get(oid, 0) for oid in all_ids if oid not in EXCLUDED_OPERATOR_IDS}
    payload["amo_meetings"] = {oid: amo_metrics["meeting"].get(oid, 0) for oid in all_ids if oid not in EXCLUDED_OPERATOR_IDS}
    payload["amo_success"] = {oid: amo_metrics["success"].get(oid, 0) for oid in all_ids if oid not in EXCLUDED_OPERATOR_IDS}
    payload["amo_revenue"] = {oid: amo_metrics["revenue"].get(oid, 0) for oid in all_ids if oid not in EXCLUDED_OPERATOR_IDS}
    payload["amo_calls_1m"] = {oid: amo_calls_1m.get(oid, 0) for oid in all_ids if oid not in EXCLUDED_OPERATOR_IDS}
    op_ids = set(payload.get("operators") or {})
    payload["inv_ck"] = {oid: (payload.get("ck_lead", {}) or {}).get(oid, 0) for oid in op_ids}
    payload["ag_ca"] = {oid: ulya_amo["ag_ca"].get(oid, 0) for oid in op_ids}
    return payload

def build_operator_name_map(operators_map):
    out = {}
    for oid, name in (operators_map or {}).items():
        out[normalize_name(short_name(name))] = oid
        out[normalize_name(name)] = oid
    return out

def ck_counts_from_sheet(date_str, operators_map):
    if not CK_SHEET_CSV_URL:
        print("CK sheet disabled: CK_SHEET_CSV_URL is empty")
        return None
    sheet = fetch_ck_sheet()
    if not sheet:
        return None
    target_date = parse_date(date_str)
    name_map = build_operator_name_map(operators_map)
    counts = Counter()
    unknown_names = set()
    for row_date, operator_name, is_ck in sheet["rows"]:
        if row_date != target_date or not is_ck:
            continue
        key = normalize_name(operator_name)
        oid = name_map.get(key)
        if not oid:
            unknown_names.add(operator_name)
            continue
        counts[oid] += 1

    if unknown_names:
        print(f"CK sheet: unknown operators: {', '.join(sorted(unknown_names))}")
    print(f"CK sheet {date_str}: rows={sum(counts.values())}")
    return counts

def fetch_ck_sheet():
    if not CK_SHEET_CSV_URL:
        print("CK sheet disabled: CK_SHEET_CSV_URL is empty")
        return None
    try:
        r = requests.get(CK_SHEET_CSV_URL, timeout=REQUEST_TIMEOUT)
    except requests.RequestException as e:
        print(f"CK sheet error: {str(e)}")
        return None
    if r.status_code != 200:
        print(f"CK sheet error: HTTP {r.status_code}")
        return None

    text = r.content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    headers = next(reader, None)
    if not headers:
        print("CK sheet error: empty header row")
        return None

    headers = [h.lstrip("\ufeff").strip() for h in headers]
    normalized = [normalize_name(h).replace(" ", "") for h in headers]

    def find_idx(options):
        for opt in options:
            key = normalize_name(opt).replace(" ", "")
            if key in normalized:
                return normalized.index(key)
        return None

    date_idx = find_idx(["Дата ЦК"])
    operator_idx = find_idx(["Передал", "Передала", "Передающий"])
    ck_idx = find_idx(["ЦК/Не ЦК", "ЦК/НЕЦК", "ЦК/НеЦК", "ЦК/НЕ ЦК"])

    if date_idx is None or operator_idx is None or ck_idx is None:
        print("CK sheet error: missing required headers (Дата ЦК, Передал, ЦК/Не ЦК)")
        print(f"CK sheet headers: {headers}")
        return None

    rows = []
    for row in reader:
        if len(row) <= max(date_idx, operator_idx, ck_idx):
            continue
        row_date = parse_sheet_date(row[date_idx])
        if not row_date:
            continue
        ck_value = normalize_name(row[ck_idx]).replace(" ", "")
        is_ck = ck_value == "цк"
        operator_name = row[operator_idx].strip()
        if not operator_name:
            continue
        rows.append((row_date, operator_name, is_ck))

    return {"rows": rows}

def sync_ck_sheet_all():
    fetch_operators()
    sheet = fetch_ck_sheet()
    if not sheet:
        return {"status": "skipped"}

    name_map = build_operator_name_map(OPERATORS)
    grouped = {}
    unknown_names = set()
    for row_date, operator_name, is_ck in sheet["rows"]:
        if not is_ck:
            continue
        key = normalize_name(operator_name)
        oid = name_map.get(key)
        if not oid:
            unknown_names.add(operator_name)
            continue
        date_str = format_date(row_date)
        grouped.setdefault(date_str, Counter())
        grouped[date_str][oid] += 1

    if unknown_names:
        print(f"CK sheet: unknown operators: {', '.join(sorted(unknown_names))}")

    conn = db_connection()
    try:
        now = datetime.now(pytz.timezone("Europe/Samara")).strftime("%Y-%m-%d %H:%M:%S")
        for date_str, counts in grouped.items():
            conn.execute("UPDATE daily_operator_stats SET ck_lead_calls = 0 WHERE date = ?", (date_str,))
            for oid, count in counts.items():
                row = conn.execute(
                    "SELECT 1 FROM daily_operator_stats WHERE date = ? AND operator_id = ?",
                    (date_str, oid)
                ).fetchone()
                if row:
                    conn.execute(
                        "UPDATE daily_operator_stats SET ck_lead_calls = ?, updated_at = ? WHERE date = ? AND operator_id = ?",
                        (count, now, date_str, oid)
                    )
                else:
                    conn.execute(
                        """
                        INSERT INTO daily_operator_stats
                        (date, operator_id, operator_name, all_calls, total_calls, cs8_calls, cs20_calls, cs22_calls, lead_agent_calls, line_calls, ck_lead_calls, talk_sum, talk_count, updated_at)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (date_str, oid, OPERATORS.get(oid, oid), 0, 0, 0, 0, 0, 0, 0, count, 0, 0, now)
                    )
        conn.commit()
    finally:
        conn.close()

    print(f"CK sheet full sync: dates={len(grouped)}")
    return {"status": "ok", "dates": len(grouped)}

def init_db():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    try:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS daily_operator_stats (
                date TEXT NOT NULL,
                operator_id TEXT NOT NULL,
                operator_name TEXT NOT NULL,
                all_calls INTEGER NOT NULL,
                total_calls INTEGER NOT NULL,
                cs8_calls INTEGER NOT NULL,
                cs20_calls INTEGER NOT NULL,
                cs22_calls INTEGER NOT NULL,
                lead_agent_calls INTEGER NOT NULL DEFAULT 0,
                ag_transfer_calls INTEGER NOT NULL DEFAULT 0,
                ag_agreement_calls INTEGER NOT NULL DEFAULT 0,
                inv_ck_calls INTEGER NOT NULL DEFAULT 0,
                ag_ca_calls INTEGER NOT NULL DEFAULT 0,
                line_calls INTEGER NOT NULL DEFAULT 0,
                ck_lead_calls INTEGER NOT NULL DEFAULT 0,
                amo_calls_1m INTEGER NOT NULL DEFAULT 0,
                amo_agreements INTEGER NOT NULL DEFAULT 0,
                amo_meetings INTEGER NOT NULL DEFAULT 0,
                amo_deals INTEGER NOT NULL DEFAULT 0,
                amo_revenue INTEGER NOT NULL DEFAULT 0,
                talk_sum INTEGER NOT NULL,
                talk_count INTEGER NOT NULL,
                updated_at TEXT NOT NULL,
                PRIMARY KEY (date, operator_id)
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS empty_days (
                date TEXT PRIMARY KEY,
                updated_at TEXT NOT NULL
            )
            """
        )
        cols = [row[1] for row in conn.execute("PRAGMA table_info(daily_operator_stats)")]
        if "lead_agent_calls" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN lead_agent_calls INTEGER NOT NULL DEFAULT 0")
        if "ag_transfer_calls" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN ag_transfer_calls INTEGER NOT NULL DEFAULT 0")
        if "ag_agreement_calls" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN ag_agreement_calls INTEGER NOT NULL DEFAULT 0")
        if "inv_ck_calls" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN inv_ck_calls INTEGER NOT NULL DEFAULT 0")
        if "ag_ca_calls" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN ag_ca_calls INTEGER NOT NULL DEFAULT 0")
        if "line_calls" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN line_calls INTEGER NOT NULL DEFAULT 0")
        if "ck_lead_calls" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN ck_lead_calls INTEGER NOT NULL DEFAULT 0")
        if "amo_calls_1m" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN amo_calls_1m INTEGER NOT NULL DEFAULT 0")
        if "amo_agreements" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN amo_agreements INTEGER NOT NULL DEFAULT 0")
        if "amo_meetings" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN amo_meetings INTEGER NOT NULL DEFAULT 0")
        if "amo_deals" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN amo_deals INTEGER NOT NULL DEFAULT 0")
        if "amo_revenue" not in cols:
            conn.execute("ALTER TABLE daily_operator_stats ADD COLUMN amo_revenue INTEGER NOT NULL DEFAULT 0")
        conn.commit()
    finally:
        conn.close()

def db_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def fetch_existing_amo_ids(date_str):
    conn = db_connection()
    try:
        rows = conn.execute(
            """
            SELECT operator_id
            FROM daily_operator_stats
            WHERE date = ?
              AND (
                amo_calls_1m > 0
                OR amo_agreements > 0
                OR amo_meetings > 0
                OR amo_deals > 0
                OR amo_revenue > 0
                OR inv_ck_calls > 0
                OR ag_ca_calls > 0
              )
            """,
            (date_str,),
        ).fetchall()
        return {row["operator_id"] for row in rows if row["operator_id"]}
    finally:
        conn.close()

def get_setting(key, default=None):
    conn = db_connection()
    try:
        row = conn.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
        return row["value"] if row else default
    finally:
        conn.close()

def set_setting(key, value):
    conn = db_connection()
    try:
        conn.execute(
            "INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (key, str(value))
        )
        conn.commit()
    finally:
        conn.close()

def parse_date(date_str):
    return datetime.strptime(date_str, "%d-%m-%Y").date()

def format_date(d):
    return d.strftime("%d-%m-%Y")

def list_saved_dates():
    conn = db_connection()
    try:
        rows = conn.execute("SELECT DISTINCT date FROM daily_operator_stats ORDER BY date").fetchall()
        return [row["date"] for row in rows]
    finally:
        conn.close()

def list_empty_dates():
    conn = db_connection()
    try:
        rows = conn.execute("SELECT date FROM empty_days ORDER BY date").fetchall()
        return [row["date"] for row in rows]
    finally:
        conn.close()

def set_empty_day(date_str, is_empty):
    conn = db_connection()
    try:
        if is_empty:
            conn.execute(
                "INSERT OR REPLACE INTO empty_days (date, updated_at) VALUES (?, ?)",
                (date_str, datetime.now(pytz.timezone("Europe/Samara")).strftime("%Y-%m-%d %H:%M:%S")),
            )
        else:
            conn.execute("DELETE FROM empty_days WHERE date = ?", (date_str,))
        conn.commit()
    finally:
        conn.close()

def ensure_range_synced(start_date, end_date):
    if not start_date or not end_date:
        return []
    start = parse_date(start_date)
    end = parse_date(end_date)
    if start > end:
        start, end = end, start
    available = set(list_saved_dates())
    missing = []
    d = start
    while d <= end:
        ds = format_date(d)
        if ds not in available:
            missing.append(ds)
        d += timedelta(days=1)
    if not missing:
        return []
    for ds in missing:
        try:
            sync_day(ds)
        except Exception as e:
            print(f"Auto sync failed for {ds}: {e}")
    return missing

def get_missing_dates_range(start_date, end_date):
    if not start_date or not end_date:
        return []
    start = parse_date(start_date)
    end = parse_date(end_date)
    if start > end:
        start, end = end, start
    available = set(list_saved_dates())
    empty = set(list_empty_dates())
    missing = []
    d = start
    while d <= end:
        ds = format_date(d)
        if ds not in available and ds not in empty:
            missing.append(ds)
        d += timedelta(days=1)
    return missing

def get_ck_lead_counts_from_db(date_str):
    conn = db_connection()
    try:
        rows = conn.execute(
            "SELECT operator_id, ck_lead_calls FROM daily_operator_stats WHERE date = ?",
            (date_str,)
        ).fetchall()
        return {row["operator_id"]: row["ck_lead_calls"] or 0 for row in rows}
    finally:
        conn.close()

def update_ck_lead_from_sheet(date_str, operators_map):
    counts = ck_counts_from_sheet(date_str, operators_map)
    if counts is None:
        return
    conn = db_connection()
    try:
        now = datetime.now(pytz.timezone("Europe/Samara")).strftime("%Y-%m-%d %H:%M:%S")
        conn.execute("UPDATE daily_operator_stats SET ck_lead_calls = 0 WHERE date = ?", (date_str,))
        for oid, count in counts.items():
            row = conn.execute(
                "SELECT 1 FROM daily_operator_stats WHERE date = ? AND operator_id = ?",
                (date_str, oid)
            ).fetchone()
            if row:
                conn.execute(
                    "UPDATE daily_operator_stats SET ck_lead_calls = ?, updated_at = ? WHERE date = ? AND operator_id = ?",
                    (count, now, date_str, oid)
                )
            else:
                conn.execute(
                    """
                    INSERT INTO daily_operator_stats
                    (date, operator_id, operator_name, all_calls, total_calls, cs8_calls, cs20_calls, cs22_calls, lead_agent_calls, line_calls, ck_lead_calls, talk_sum, talk_count, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (date_str, oid, operators_map.get(oid, oid), 0, 0, 0, 0, 0, 0, 0, count, 0, 0, now)
                )
        conn.commit()
    finally:
        conn.close()

def has_date_in_db(date_str):
    conn = db_connection()
    try:
        row = conn.execute(
            "SELECT 1 FROM daily_operator_stats WHERE date = ? LIMIT 1",
            (date_str,)
        ).fetchone()
        return row is not None
    finally:
        conn.close()

def has_call_data_for_date(date_str):
    conn = db_connection()
    try:
        row = conn.execute(
            "SELECT 1 FROM daily_operator_stats WHERE date = ? AND all_calls > 0 LIMIT 1",
            (date_str,)
        ).fetchone()
        return row is not None
    finally:
        conn.close()

def has_nonzero_stats_for_date(date_str):
    conn = db_connection()
    try:
        row = conn.execute(
            """
            SELECT 1
            FROM daily_operator_stats
            WHERE date = ?
              AND (
                all_calls > 0 OR total_calls > 0 OR cs8_calls > 0 OR cs20_calls > 0 OR cs22_calls > 0
                OR lead_agent_calls > 0 OR ag_transfer_calls > 0 OR ag_agreement_calls > 0
                OR inv_ck_calls > 0 OR ag_ca_calls > 0 OR line_calls > 0 OR ck_lead_calls > 0
                OR amo_calls_1m > 0 OR amo_agreements > 0 OR amo_meetings > 0 OR amo_deals > 0 OR amo_revenue > 0
              )
            LIMIT 1
            """,
            (date_str,)
        ).fetchone()
        return row is not None
    finally:
        conn.close()

def get_day_stats_from_db(date_str):
    moscow_ids = list(MOSCOW_OPERATOR_IDS) or ["__none__"]
    conn = db_connection()
    try:
        rows = conn.execute(
            """
            SELECT
                operator_id,
                operator_name,
                all_calls,
                total_calls,
                cs8_calls,
                cs20_calls,
                ag_transfer_calls,
                ag_agreement_calls,
                inv_ck_calls,
                ag_ca_calls,
                lead_agent_calls,
                line_calls,
                ck_lead_calls,
                amo_calls_1m,
                amo_agreements,
                amo_meetings,
                amo_deals,
                amo_revenue,
                talk_sum,
                talk_count
            FROM daily_operator_stats
            WHERE date = ?
            """,
            (date_str,)
        ).fetchall()
    finally:
        conn.close()

    if not rows:
        return None

    by_id = {row["operator_id"]: row for row in rows if row["operator_id"] not in EXCLUDED_OPERATOR_IDS}
    active_ids = {
        oid for oid, row in by_id.items()
        if (row["all_calls"] or 0) > 0
        or (row["inv_ck_calls"] or 0) > 0
        or (row["ag_ca_calls"] or 0) > 0
    }
    amo_ids = {
        oid for oid, row in by_id.items()
        if (row["amo_calls_1m"] or 0) > 0
        or (row["amo_agreements"] or 0) > 0
        or (row["amo_meetings"] or 0) > 0
        or (row["amo_deals"] or 0) > 0
        or (row["amo_revenue"] or 0) > 0
        or oid in MOSCOW_OPERATOR_IDS
    }
    amo_ids = {oid for oid in amo_ids if oid not in EXCLUDED_OPERATOR_IDS}
    if MOSCOW_OPERATOR_NAME_ALLOWLIST:
        amo_ids = {
            oid for oid in amo_ids
            if moscow_name_allowed(short_name(by_id[oid]["operator_name"] or oid))
        }
    operators = {oid: short_name(by_id[oid]["operator_name"] or oid) for oid in active_ids}
    amo_operators = {oid: short_name(by_id[oid]["operator_name"] or oid) for oid in amo_ids}

    def avg_for(row):
        talk_count = row["talk_count"] or 0
        if talk_count <= 0:
            return 0
        return int((row["talk_sum"] or 0) / talk_count)

    return {
        "operators": operators,
        "amo_operators": amo_operators,
        "status": {oid: "" for oid in operators},
        "all": {oid: by_id[oid]["all_calls"] or 0 for oid in operators},
        "line": {oid: format_hms(by_id[oid]["line_calls"] or 0) for oid in operators},
        "total": {oid: by_id[oid]["total_calls"] or 0 for oid in operators},
        "cs8": {oid: by_id[oid]["cs8_calls"] or 0 for oid in operators},
        "cs20": {oid: by_id[oid]["cs20_calls"] or 0 for oid in operators},
        "ag_transfer": {oid: by_id[oid]["ag_transfer_calls"] or 0 for oid in operators},
        "ag_agreement": {oid: by_id[oid]["ag_agreement_calls"] or 0 for oid in operators},
        "inv_ck": {oid: by_id[oid]["inv_ck_calls"] or 0 for oid in operators},
        "ag_ca": {oid: by_id[oid]["ag_ca_calls"] or 0 for oid in operators},
        "lead_agent": {oid: by_id[oid]["lead_agent_calls"] or 0 for oid in operators},
        "ck_lead": {oid: by_id[oid]["ck_lead_calls"] or 0 for oid in operators},
        "avg": {oid: avg_for(by_id[oid]) for oid in operators},
        "amo_calls_1m": {oid: by_id[oid]["amo_calls_1m"] or 0 for oid in amo_operators},
        "amo_agreements": {oid: by_id[oid]["amo_agreements"] or 0 for oid in amo_operators},
        "amo_meetings": {oid: by_id[oid]["amo_meetings"] or 0 for oid in amo_operators},
        "amo_deals": {oid: by_id[oid]["amo_deals"] or 0 for oid in amo_operators},
        "amo_revenue": {oid: by_id[oid]["amo_revenue"] or 0 for oid in amo_operators},
        "new": 0,
        "new_noactive": 0,
        "server_time": int(time.time() * 1000)
    }

def to_db_key(date_str):
    return parse_date(date_str).strftime("%Y%m%d")

def get_db_range():
    conn = db_connection()
    try:
        row = conn.execute(
            """
            SELECT date FROM daily_operator_stats
            ORDER BY (substr(date,7,4)||substr(date,4,2)||substr(date,1,2)) ASC
            LIMIT 1
            """
        ).fetchone()
        row_max = conn.execute(
            """
            SELECT date FROM daily_operator_stats
            ORDER BY (substr(date,7,4)||substr(date,4,2)||substr(date,1,2)) DESC
            LIMIT 1
            """
        ).fetchone()
        return (row["date"] if row else None, row_max["date"] if row_max else None)
    finally:
        conn.close()

def get_report_data(start_date=None, end_date=None):
    db_start, db_end = get_db_range()
    if not db_start or not db_end:
        return {
            "range": None,
            "rows": [],
            "totals": {}
        }
    if not start_date or not end_date:
        start_date, end_date = db_start, db_end
    if parse_date(start_date) > parse_date(end_date):
        start_date, end_date = end_date, start_date

    start_key = to_db_key(start_date)
    end_key = to_db_key(end_date)

    conn = db_connection()
    try:
        rows = conn.execute(
            """
            SELECT
                operator_id,
                operator_name,
                SUM(all_calls) AS all_calls,
                SUM(total_calls) AS total_calls,
                SUM(cs8_calls) AS cs8_calls,
                SUM(cs20_calls) AS cs20_calls,
                SUM(cs22_calls) AS cs22_calls,
                SUM(lead_agent_calls) AS lead_agent_calls,
                SUM(ag_transfer_calls) AS ag_transfer_calls,
                SUM(ag_agreement_calls) AS ag_agreement_calls,
                SUM(inv_ck_calls) AS inv_ck_calls,
                SUM(ag_ca_calls) AS ag_ca_calls,
                SUM(line_calls) AS line_calls,
                SUM(ck_lead_calls) AS ck_lead_calls,
                SUM(talk_sum) AS talk_sum,
                SUM(talk_count) AS talk_count
            FROM daily_operator_stats
            WHERE (substr(date,7,4)||substr(date,4,2)||substr(date,1,2)) BETWEEN ? AND ?
            GROUP BY operator_id, operator_name
            ORDER BY operator_name
            """,
            (start_key, end_key)
        ).fetchall()
    finally:
        conn.close()

    result_rows = []
    totals = {
        "all": 0,
        "total": 0,
        "cs8": 0,
        "cs20": 0,
        "cs22": 0,
        "lead_agent": 0,
        "ag_transfer": 0,
        "ag_agreement": 0,
        "inv_ck": 0,
        "ag_ca": 0,
        "line": 0,
        "ck_lead": 0,
        "talk_sum": 0,
        "talk_count": 0
    }
    for row in rows:
        if row["operator_id"] in EXCLUDED_OPERATOR_IDS:
            continue
        talk_count = row["talk_count"] or 0
        talk_sum = row["talk_sum"] or 0
        avg = int(talk_sum / talk_count) if talk_count else 0
        result_rows.append({
            "operator_id": row["operator_id"],
            "operator_name": row["operator_name"],
            "all": row["all_calls"] or 0,
            "total": row["total_calls"] or 0,
            "cs8": row["cs8_calls"] or 0,
            "cs20": row["cs20_calls"] or 0,
            "cs22": row["cs22_calls"] or 0,
            "lead_agent": row["lead_agent_calls"] or 0,
            "ag_transfer": row["ag_transfer_calls"] or 0,
            "ag_agreement": row["ag_agreement_calls"] or 0,
            "inv_ck": row["inv_ck_calls"] or 0,
            "ag_ca": row["ag_ca_calls"] or 0,
            "line": row["line_calls"] or 0,
            "ck_lead": row["ck_lead_calls"] or 0,
            "avg": avg,
            "talk_sum": talk_sum,
            "talk_count": talk_count
        })
        totals["all"] += row["all_calls"] or 0
        totals["total"] += row["total_calls"] or 0
        totals["cs8"] += row["cs8_calls"] or 0
        totals["cs20"] += row["cs20_calls"] or 0
        totals["cs22"] += row["cs22_calls"] or 0
        totals["lead_agent"] += row["lead_agent_calls"] or 0
        totals["ag_transfer"] += row["ag_transfer_calls"] or 0
        totals["ag_agreement"] += row["ag_agreement_calls"] or 0
        totals["inv_ck"] += row["inv_ck_calls"] or 0
        totals["ag_ca"] += row["ag_ca_calls"] or 0
        totals["line"] += row["line_calls"] or 0
        totals["ck_lead"] += row["ck_lead_calls"] or 0
        totals["talk_sum"] += talk_sum
        totals["talk_count"] += talk_count

    avg_total = int(totals["talk_sum"] / totals["talk_count"]) if totals["talk_count"] else 0
    reach = round((totals["total"] / totals["all"]) * 100, 1) if totals["all"] else 0

    return {
        "range": {"start": start_date, "end": end_date},
        "rows": result_rows,
        "totals": {
            "all": totals["all"],
            "total": totals["total"],
            "cs8": totals["cs8"],
            "cs20": totals["cs20"],
            "cs22": totals["cs22"],
            "lead_agent": totals["lead_agent"],
            "ag_transfer": totals["ag_transfer"],
            "ag_agreement": totals["ag_agreement"],
            "inv_ck": totals["inv_ck"],
            "ag_ca": totals["ag_ca"],
            "line": totals["line"],
            "ck_lead": totals["ck_lead"],
            "avg": avg_total,
            "reach": reach
        }
    }

def get_moscow_report_data_db(start_date=None, end_date=None):
    db_start, db_end = get_db_range()
    if not db_start or not db_end:
        return {"range": None, "rows": [], "totals": {}}
    if not start_date or not end_date:
        start_date, end_date = db_start, db_end
    if parse_date(start_date) > parse_date(end_date):
        start_date, end_date = end_date, start_date

    start_key = to_db_key(start_date)
    end_key = to_db_key(end_date)
    moscow_ids = list(MOSCOW_OPERATOR_IDS) or ["__none__"]
    conn = db_connection()
    try:
        rows = conn.execute(
            """
            SELECT
                operator_id,
                operator_name,
                SUM(amo_calls_1m) AS amo_calls_1m,
                SUM(amo_agreements) AS amo_agreements,
                SUM(amo_meetings) AS amo_meetings,
                SUM(amo_deals) AS amo_deals,
                SUM(amo_revenue) AS amo_revenue
            FROM daily_operator_stats
            WHERE (substr(date,7,4)||substr(date,4,2)||substr(date,1,2)) BETWEEN ? AND ?
              AND (
                amo_calls_1m > 0 OR amo_agreements > 0 OR amo_meetings > 0 OR amo_deals > 0 OR amo_revenue > 0
                OR operator_id IN ({placeholders})
              )
            GROUP BY operator_id, operator_name
            ORDER BY operator_name
            """.format(placeholders=",".join(["?"] * len(moscow_ids))),
            (start_key, end_key, *moscow_ids)
        ).fetchall()
    finally:
        conn.close()

    merged = {}
    totals = {
        "calls_1m": 0,
        "agreements": 0,
        "meetings": 0,
        "deals": 0,
        "revenue": 0
    }
    for row in rows:
        if row["operator_id"] in EXCLUDED_OPERATOR_IDS:
            continue
        if not moscow_name_allowed(row["operator_name"]):
            continue
        name = row["operator_name"]
        data = {
            "operator_id": row["operator_id"],
            "operator_name": name,
            "calls_1m": row["amo_calls_1m"] or 0,
            "agreements": row["amo_agreements"] or 0,
            "meetings": row["amo_meetings"] or 0,
            "deals": row["amo_deals"] or 0,
            "revenue": row["amo_revenue"] or 0
        }
        if name in merged:
            merged[name]["calls_1m"] += data["calls_1m"]
            merged[name]["agreements"] += data["agreements"]
            merged[name]["meetings"] += data["meetings"]
            merged[name]["deals"] += data["deals"]
            merged[name]["revenue"] += data["revenue"]
        else:
            merged[name] = data

    result_rows = sorted(merged.values(), key=lambda r: r["operator_name"])
    for row in result_rows:
        totals["calls_1m"] += row["calls_1m"]
        totals["agreements"] += row["agreements"]
        totals["meetings"] += row["meetings"]
        totals["deals"] += row["deals"]
        totals["revenue"] += row["revenue"]

    return {
        "range": {"start": start_date, "end": end_date},
        "rows": result_rows,
        "totals": totals
    }

def amo_report_data(start_date=None, end_date=None):
    if not start_date or not end_date:
        return {"range": None, "rows": [], "totals": {}}
    if parse_date(start_date) > parse_date(end_date):
        start_date, end_date = end_date, start_date
    metrics = amo_leads_created_metrics_range(start_date, end_date)
    calls_1m = amo_calls_over_minute_range(start_date, end_date)
    amo_users = amo_fetch_users()
    all_ids = set(metrics["agreement"]) | set(metrics["meeting"]) | set(metrics["success"]) | set(metrics["revenue"]) | set(calls_1m)
    rows = []
    totals = {
        "calls_1m": 0,
        "agreements": 0,
        "meetings": 0,
        "deals": 0,
        "revenue": 0
    }
    for oid in all_ids:
        name = short_name(amo_users.get(oid, f"User {oid}"))
        if not moscow_name_allowed(name):
            continue
        row = {
            "operator_id": oid,
            "operator_name": name,
            "calls_1m": calls_1m.get(oid, 0),
            "agreements": metrics["agreement"].get(oid, 0),
            "meetings": metrics["meeting"].get(oid, 0),
            "deals": metrics["success"].get(oid, 0),
            "revenue": metrics["revenue"].get(oid, 0)
        }
        rows.append(row)
        totals["calls_1m"] += row["calls_1m"]
        totals["agreements"] += row["agreements"]
        totals["meetings"] += row["meetings"]
        totals["deals"] += row["deals"]
        totals["revenue"] += row["revenue"]

    if not MOSCOW_OPERATOR_NAME_ALLOWLIST:
        amo_name_set = {short_name(name) for name in amo_users.values()}
        for oid in MOSCOW_OPERATOR_IDS:
            if oid in all_ids:
                continue
            name = short_name(OPERATORS.get(oid) or f"User {oid}")
            if name in amo_name_set:
                continue
            rows.append({
                "operator_id": oid,
                "operator_name": name,
                "calls_1m": 0,
                "agreements": 0,
                "meetings": 0,
                "deals": 0,
                "revenue": 0
            })

    rows.sort(key=lambda r: r["operator_name"])
    return {
        "range": {"start": start_date, "end": end_date},
        "rows": rows,
        "totals": totals
    }

def filter_report_rows(rows, branch):
    if not branch:
        return rows
    branch = branch.lower()
    if branch == "moscow":
        if MOSCOW_OPERATOR_NAME_ALLOWLIST:
            return [row for row in rows if moscow_name_allowed(row.get("operator_name", ""))]
        return [row for row in rows if row["operator_id"] in MOSCOW_OPERATOR_IDS]
    if branch == "ulyanovsk":
        filtered = []
        for row in rows:
            if row["operator_id"] in MOSCOW_OPERATOR_IDS:
                continue
            sip_total = (
                row.get("all", 0)
                + row.get("total", 0)
                + row.get("cs8", 0)
                + row.get("cs20", 0)
                + row.get("cs22", 0)
                + row.get("lead_agent", 0)
                + row.get("ag_transfer", 0)
                + row.get("ag_agreement", 0)
                + row.get("inv_ck", 0)
                + row.get("ag_ca", 0)
                + row.get("line", 0)
                + row.get("ck_lead", 0)
            )
            if sip_total == 0:
                continue
            filtered.append(row)
        return filtered
    return rows

def recalc_report_totals(rows):
    totals = {
        "all": 0,
        "total": 0,
        "cs8": 0,
        "cs20": 0,
        "cs22": 0,
        "lead_agent": 0,
        "ag_transfer": 0,
        "ag_agreement": 0,
        "inv_ck": 0,
        "ag_ca": 0,
        "line": 0,
        "ck_lead": 0,
        "talk_sum": 0,
        "talk_count": 0
    }
    for row in rows:
        totals["all"] += row.get("all", 0)
        totals["total"] += row.get("total", 0)
        totals["cs8"] += row.get("cs8", 0)
        totals["cs20"] += row.get("cs20", 0)
        totals["cs22"] += row.get("cs22", 0)
        totals["lead_agent"] += row.get("lead_agent", 0)
        totals["ag_transfer"] += row.get("ag_transfer", 0)
        totals["ag_agreement"] += row.get("ag_agreement", 0)
        totals["inv_ck"] += row.get("inv_ck", 0)
        totals["ag_ca"] += row.get("ag_ca", 0)
        totals["line"] += row.get("line", 0)
        totals["ck_lead"] += row.get("ck_lead", 0)
        totals["talk_sum"] += row.get("talk_sum", 0)
        totals["talk_count"] += row.get("talk_count", 0)
    avg_total = int(totals["talk_sum"] / totals["talk_count"]) if totals["talk_count"] else 0
    reach = round((totals["total"] / totals["all"]) * 100, 1) if totals["all"] else 0
    return {
        "all": totals["all"],
        "total": totals["total"],
        "cs8": totals["cs8"],
        "cs20": totals["cs20"],
        "cs22": totals["cs22"],
        "lead_agent": totals["lead_agent"],
        "ag_transfer": totals["ag_transfer"],
        "ag_agreement": totals["ag_agreement"],
        "inv_ck": totals["inv_ck"],
        "ag_ca": totals["ag_ca"],
        "line": totals["line"],
        "ck_lead": totals["ck_lead"],
        "avg": avg_total,
        "reach": reach
    }
def fetch_calls_for_date(date_str, operators_map=None):
    params_base = build_base_params(requested_date=date_str, operators_map=operators_map)
    page = 1
    limit = 1000
    items = []
    while True:
        params = params_base + [("page", page), ("limit", limit)]
        r = requests.get(CALL_LIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT)
        r.raise_for_status()
        chunk = r.json().get("items", [])
        if not chunk:
            break
        items.extend(chunk)
        if len(chunk) < limit:
            break
        page += 1
    return items

def aggregate_calls(calls, operators_map=None):
    stats = defaultdict(lambda: {
        "all": 0,
        "total": 0,
        "cs8": 0,
        "cs20": 0,
        "cs22": 0,
        "lead_agent": 0,
        "ag_transfer": 0,
        "ag_agreement": 0,
        "line": 0,
        "ck_lead": 0,
        "amo_calls_1m": 0,
        "amo_agreements": 0,
        "amo_meetings": 0,
        "amo_deals": 0,
        "amo_revenue": 0,
        "talk_sum": 0,
        "talk_count": 0,
        "name": ""
    })
    for call in calls:
        op = call.get("operator") or {}
        oid = str(op.get("id") or "")
        if not oid:
            continue
        if operators_map and oid not in operators_map:
            continue
        name = op.get("full_name") or op.get("fullName") or operators_map.get(oid, "") if operators_map else ""
        status = get_status_id(call)
        td = get_talk_duration(call)
        is_dialog = status in STAT_FULL and td >= 20
        stats[oid]["all"] += 1
        if is_dialog:
            stats[oid]["total"] += 1
            stats[oid]["talk_sum"] += td
            stats[oid]["talk_count"] += 1
        if status in CS8:
            stats[oid]["cs8"] += 1
        if status in CS20:
            stats[oid]["cs20"] += 1
        if status in CS22:
            stats[oid]["cs22"] += 1
        if status in LEAD_AGENT:
            stats[oid]["lead_agent"] += 1
        if status in AG_TRANSFER:
            stats[oid]["ag_transfer"] += 1
        if status in AG_AGREEMENT:
            stats[oid]["ag_agreement"] += 1
        if name:
            stats[oid]["name"] = name
    return stats

def upsert_daily_stats(date_str, stats):
    now = datetime.now(pytz.timezone("Europe/Samara")).strftime("%Y-%m-%d %H:%M:%S")
    conn = db_connection()
    try:
        for oid, data in stats.items():
            name = data["name"] or oid
            conn.execute(
                """
                INSERT INTO daily_operator_stats
                (date, operator_id, operator_name, all_calls, total_calls, cs8_calls, cs20_calls, cs22_calls, lead_agent_calls, ag_transfer_calls, ag_agreement_calls, inv_ck_calls, ag_ca_calls, line_calls, ck_lead_calls, amo_calls_1m, amo_agreements, amo_meetings, amo_deals, amo_revenue, talk_sum, talk_count, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(date, operator_id) DO UPDATE SET
                    operator_name=excluded.operator_name,
                    all_calls=excluded.all_calls,
                    total_calls=excluded.total_calls,
                    cs8_calls=excluded.cs8_calls,
                    cs20_calls=excluded.cs20_calls,
                    cs22_calls=excluded.cs22_calls,
                    lead_agent_calls=excluded.lead_agent_calls,
                    ag_transfer_calls=excluded.ag_transfer_calls,
                    ag_agreement_calls=excluded.ag_agreement_calls,
                    inv_ck_calls=excluded.inv_ck_calls,
                    ag_ca_calls=excluded.ag_ca_calls,
                    line_calls=excluded.line_calls,
                    ck_lead_calls=excluded.ck_lead_calls,
                    amo_calls_1m=excluded.amo_calls_1m,
                    amo_agreements=excluded.amo_agreements,
                    amo_meetings=excluded.amo_meetings,
                    amo_deals=excluded.amo_deals,
                    amo_revenue=excluded.amo_revenue,
                    talk_sum=excluded.talk_sum,
                    talk_count=excluded.talk_count,
                    updated_at=excluded.updated_at
                """,
                (
                    date_str,
                    oid,
                    name,
                    data["all"],
                    data["total"],
                    data["cs8"],
                    data["cs20"],
                    data["cs22"],
                    data["lead_agent"],
                    data.get("ag_transfer", 0),
                    data.get("ag_agreement", 0),
                    data.get("inv_ck", 0),
                    data.get("ag_ca", 0),
                    data["line"],
                    data["ck_lead"],
                    data.get("amo_calls_1m", 0),
                    data.get("amo_agreements", 0),
                    data.get("amo_meetings", 0),
                    data.get("amo_deals", 0),
                    data.get("amo_revenue", 0),
                    data["talk_sum"],
                    data["talk_count"],
                    now
                )
            )
        conn.commit()
    finally:
        conn.close()

def sync_day(date_str):
    sip_ok = True
    ck_ok = True
    amo_ok = True
    try:
        fetch_operators()
    except Exception as e:
        print(f"SipSpeak operators fetch failed: {e}")
        sip_ok = False
    try:
        calls = fetch_calls_for_date(date_str, operators_map=OPERATORS)
    except Exception as e:
        print(f"SipSpeak calls fetch failed for {date_str}: {e}")
        calls = []
        sip_ok = False
    operators_from_calls = extract_operators_from_calls(calls)
    operators_map = OPERATORS or operators_from_calls
    try:
        line_seconds = fetch_line_seconds(date_str, operators_map=operators_map)
    except Exception as e:
        print(f"SipSpeak line fetch failed for {date_str}: {e}")
        line_seconds = {}
        sip_ok = False
    try:
        ck_counts = ck_counts_from_sheet(date_str, operators_map)
    except Exception as e:
        print(f"CK sheet fetch failed for {date_str}: {e}")
        ck_counts = {}
        ck_ok = False
    stats = aggregate_calls(calls, operators_map=operators_map)
    # Надежный пересчет агентских метрик отдельными фильтрованными запросами.
    # Это защищает от случаев, когда в полном списке звонков формат статуса отличается.
    try:
        ag_transfer_counts = fetch_counts(AG_TRANSFER, requested_date=date_str, operators_map=operators_map)
        ag_agreement_counts = fetch_counts(AG_AGREEMENT, requested_date=date_str, operators_map=operators_map)
    except Exception as e:
        print(f"SipSpeak agent status fetch failed for {date_str}: {e}")
        ag_transfer_counts = Counter()
        ag_agreement_counts = Counter()
    if ag_transfer_counts or ag_agreement_counts:
        for oid in (set(stats) | set(ag_transfer_counts) | set(ag_agreement_counts)):
            if oid not in stats:
                stats[oid] = {
                    "all": 0,
                    "total": 0,
                    "cs8": 0,
                    "cs20": 0,
                    "cs22": 0,
                    "lead_agent": 0,
                    "ag_transfer": 0,
                    "ag_agreement": 0,
                    "line": 0,
                    "ck_lead": 0,
                    "amo_calls_1m": 0,
                    "amo_agreements": 0,
                    "amo_meetings": 0,
                    "amo_deals": 0,
                    "amo_revenue": 0,
                    "talk_sum": 0,
                    "talk_count": 0,
                    "name": operators_map.get(oid, "")
                }
            stats[oid]["ag_transfer"] = ag_transfer_counts.get(oid, 0)
            stats[oid]["ag_agreement"] = ag_agreement_counts.get(oid, 0)
    if not sip_ok and not calls and has_date_in_db(date_str):
        cached = get_day_stats_from_db(date_str)
        print(f"Sync skipped for {date_str}: SipSpeak unavailable, keeping existing DB data")
        return {
            "date": date_str,
            "operators": len(cached.get("operators", {})) if cached else 0,
            "calls": 0
        }
    if ck_counts is None:
        ck_counts = {}
    for oid, count in ck_counts.items():
        if oid not in stats:
            stats[oid] = {
                "all": 0,
                "total": 0,
                "cs8": 0,
                "cs20": 0,
                "cs22": 0,
                "lead_agent": 0,
                "ag_transfer": 0,
                "ag_agreement": 0,
                "ag_ca": 0,
                "line": 0,
                "ck_lead": count,
                "inv_ck": count,
                "amo_calls_1m": 0,
                "amo_agreements": 0,
                "amo_meetings": 0,
                "amo_deals": 0,
                "amo_revenue": 0,
                "talk_sum": 0,
                "talk_count": 0,
                "name": operators_map.get(oid, "")
            }
        else:
            stats[oid]["ck_lead"] = count
            stats[oid]["inv_ck"] = count
    for oid, seconds in line_seconds.items():
        stats[oid]["line"] = seconds
        if not stats[oid].get("name"):
            stats[oid]["name"] = operators_map.get(oid, "")

    if amo_enabled():
        try:
            amo_metrics = amo_leads_event_metrics(date_str)
            amo_calls_1m = amo_calls_over_minute(date_str)
            ulya_amo = amo_ulya_invest_agent_metrics(date_str, operators_override=operators_map)
            amo_users = amo_fetch_users()
        except Exception as e:
            print(f"AMO sync failed for {date_str}: {e}")
            amo_metrics = {"agreement": Counter(), "meeting": Counter(), "success": Counter(), "revenue": defaultdict(int)}
            amo_calls_1m = Counter()
            ulya_amo = {"inv_ck": Counter(), "ag_ca": Counter()}
            amo_users = {}
            amo_ok = False
        existing_amo_ids = fetch_existing_amo_ids(date_str)
        amo_ids = (
            set(amo_metrics["agreement"])
            | set(amo_metrics["meeting"])
            | set(amo_metrics["success"])
            | set(amo_metrics["revenue"])
            | set(amo_calls_1m)
            | set(ulya_amo["ag_ca"])
            | set(MOSCOW_OPERATOR_IDS)
            | set(existing_amo_ids)
        )
        amo_ids = {oid for oid in amo_ids if oid not in EXCLUDED_OPERATOR_IDS}
        for oid in amo_ids:
            if oid not in stats:
                stats[oid] = {
                    "all": 0,
                    "total": 0,
                    "cs8": 0,
                    "cs20": 0,
                    "cs22": 0,
                    "lead_agent": 0,
                    "ag_transfer": 0,
                    "ag_agreement": 0,
                    "inv_ck": 0,
                    "ag_ca": 0,
                    "line": 0,
                    "ck_lead": 0,
                    "amo_calls_1m": 0,
                    "amo_agreements": 0,
                    "amo_meetings": 0,
                    "amo_deals": 0,
                    "amo_revenue": 0,
                    "talk_sum": 0,
                    "talk_count": 0,
                    "name": ""
                }
            stats[oid]["amo_calls_1m"] = amo_calls_1m.get(oid, 0)
            stats[oid]["amo_agreements"] = amo_metrics["agreement"].get(oid, 0)
            stats[oid]["amo_meetings"] = amo_metrics["meeting"].get(oid, 0)
            stats[oid]["amo_deals"] = amo_metrics["success"].get(oid, 0)
            stats[oid]["amo_revenue"] = amo_metrics["revenue"].get(oid, 0)
            stats[oid]["ag_ca"] = ulya_amo["ag_ca"].get(oid, stats[oid].get("ag_ca", 0))
            if not stats[oid].get("name"):
                stats[oid]["name"] = short_name(amo_users.get(oid, operators_map.get(oid, "")))
        for oid in MOSCOW_SIPSPEAK_AGREEMENTS_IDS:
            if oid in stats:
                stats[oid]["amo_agreements"] = stats[oid].get("cs8", 0)

    for oid in list(stats.keys()):
        if oid in EXCLUDED_OPERATOR_IDS:
            del stats[oid]

    for oid in stats:
        stats[oid].setdefault("amo_calls_1m", 0)
        stats[oid].setdefault("amo_agreements", 0)
        stats[oid].setdefault("amo_meetings", 0)
        stats[oid].setdefault("amo_deals", 0)
        stats[oid].setdefault("amo_revenue", 0)
        stats[oid].setdefault("ag_transfer", 0)
        stats[oid].setdefault("ag_agreement", 0)
        stats[oid].setdefault("inv_ck", 0)
        stats[oid].setdefault("ag_ca", 0)
    total_activity = 0
    for item in stats.values():
        total_activity += (
            item.get("all", 0)
            + item.get("total", 0)
            + item.get("cs8", 0)
            + item.get("cs20", 0)
            + item.get("cs22", 0)
            + item.get("lead_agent", 0)
            + item.get("ag_transfer", 0)
            + item.get("ag_agreement", 0)
            + item.get("inv_ck", 0)
            + item.get("ag_ca", 0)
            + item.get("line", 0)
            + item.get("ck_lead", 0)
            + item.get("amo_calls_1m", 0)
            + item.get("amo_agreements", 0)
            + item.get("amo_meetings", 0)
            + item.get("amo_deals", 0)
            + item.get("amo_revenue", 0)
        )
    empty_ok = (sip_ok and ck_ok and (amo_ok or not amo_enabled()) and total_activity == 0)
    set_empty_day(date_str, empty_ok)
    upsert_daily_stats(date_str, stats)
    return {
        "date": date_str,
        "operators": len(stats),
        "calls": sum(item["all"] for item in stats.values())
    }

def sync_range(start_str, end_str):
    start = parse_date(start_str)
    end = parse_date(end_str)
    if start > end:
        start, end = end, start
    out = []
    d = start
    while d <= end:
        out.append(sync_day(format_date(d)))
        d += timedelta(days=1)
    return out

def fetch_operators():
    """Получает список операторов из API"""
    global OPERATORS
    operators = {}
    page = 1
    limit = 500
    try:
        while True:
            params = [("page", page), ("limit", limit), ("removed", "false")]
            r = requests.get(USERS_LIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT)
            r.raise_for_status()
            items = r.json().get("items", [])
            if not items:
                break
            for user in items:
                if user.get("role") == "ROLE_OPERATOR":
                    uid = str(user.get("id") or "")
                    name = user.get("full_name") or user.get("fullName") or ""
                    if uid and name and uid not in EXCLUDED_OPERATOR_IDS:
                        operators[uid] = name
            if len(items) < limit:
                break
            page += 1
    except Exception as e:
        print(f"Ошибка при получении списка операторов: {str(e)}")
        return OPERATORS
    if operators:
        OPERATORS = operators
    return OPERATORS

def fetch_active_campaigns():
    """Получает список активных проектов"""
    global active_campaigns
    if not USE_ACTIVE_CAMPAIGNS:
        active_campaigns = []
        return []
    try:
        print("Получение списка активных проектов...")
        params = [("active", "true")]
        r = requests.get(CAMPAIGN_LIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT)
        r.raise_for_status()
        campaigns = r.json().get("items", [])
        active_campaigns = [str(camp.get("id")) for camp in campaigns if camp.get("id")]
        print(f"Получено {len(active_campaigns)} активных проектов: {active_campaigns}")
        return active_campaigns
    except Exception as e:
        print(f"Ошибка при получении активных проектов: {str(e)}")
        return []

def fetch_new_numbers_total_by_active():
    """Получает количество новых номеров для активных проектов"""
    if USE_ACTIVE_CAMPAIGNS:
        if not active_campaigns:
            fetch_active_campaigns()
        params = [("statuses[]", "1")]
        for campaign_id in active_campaigns:
            params.append(("campaign_ids[]", campaign_id))
    else:
        params = [("statuses[]", "1")]
    params.append(("page", 1))
    params.append(("limit", 1))

    try:
        r = requests.get(CONTACT_LIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT)
        r.raise_for_status()
        data = r.json()
        return data.get("totalCount", len(data.get("items", [])))
    except Exception as e:
        print(f"Ошибка при получении количества новых номеров: {str(e)}")
        return 0

def build_base_params(requested_date=None, operators_map=None):
    if requested_date:
        date = requested_date
    elif TEST_DATE:
        date = TEST_DATE
    else:
        tz   = pytz.timezone("Europe/Samara")
        date = datetime.now(tz).strftime("%d-%m-%Y")
    params = [
        ("start_at", f"{date} 00:00"),
        ("end_at",   f"{date} 23:59")
    ]
    if operators_map is None:
        operators_map = OPERATORS
    if operators_map:
        for op in operators_map:
            params.append(("operators[]", op))
    return params


def fetch_counts(status_list, requested_date=None, operators_map=None):
    if operators_map is None:
        operators_map = OPERATORS
    params = build_base_params(requested_date=requested_date, operators_map=operators_map) + ALL_CALLS_PARAMS
    for s in status_list:
        params.append(("client_statuses[]", s))
    r = requests.get(CALL_LIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT)
    r.raise_for_status()
    cnt = Counter()
    for it in r.json().get("items", []):
        oid = str(it.get("operator",{}).get("id") or "")
        if not operators_map or oid in operators_map:
            cnt[oid] += 1
    return cnt


def fetch_all_counts(requested_date=None, operators_map=None):
    return fetch_counts([], requested_date=requested_date, operators_map=operators_map)


def fetch_all_calls_details(requested_date=None, operators_map=None):
    params = build_base_params(requested_date=requested_date, operators_map=operators_map) + ALL_CALLS_PARAMS
    r = requests.get(CALL_LIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT)
    r.raise_for_status()
    return r.json().get("items", [])


def fetch_current_status(requested_date=None, operators_map=None):
    if operators_map is None:
        operators_map = OPERATORS
    params = build_base_params(requested_date=requested_date, operators_map=operators_map) + [("page",1),("limit",1000)]
    r1 = requests.get(HIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT); r1.raise_for_status()
    r2 = requests.get(LIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT); r2.raise_for_status()
    status = {}
    for ev in r1.json().get("items", []):
        oid = str(ev.get("id") or "")
        if (not operators_map or oid in operators_map) and ev.get("event"):
            status[oid] = ev["event"]
    for ev in r2.json().get("items", []):
        oid = str(ev.get("id") or "")
        if (not operators_map or oid in operators_map) and ev.get("event"):
            status[oid] = ev["event"]
    return {oid: STATUS_MAP.get(st, st) for oid,st in status.items()}

def fetch_line_times(requested_date=None, operators_map=None):
    if operators_map is None:
        operators_map = OPERATORS
    params = build_base_params(requested_date=requested_date, operators_map=operators_map) + [("page",1),("limit",1000)]
    r = requests.get(LIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT)
    r.raise_for_status()
    out = {}
    for item in r.json().get("items", []):
        oid = str(item.get("id") or "")
        if operators_map and oid not in operators_map:
            continue
        def to_int(value):
            try:
                return int(value or 0)
            except (TypeError, ValueError):
                return 0
        active = to_int(item.get("active"))
        dnd = to_int(item.get("dnd"))
        call = to_int(item.get("call"))
        ringing = to_int(item.get("ringing"))
        line_total = active + dnd + call + ringing
        out[oid] = format_hms(line_total)
    return out

def fetch_line_seconds(date_str, operators_map=None):
    if operators_map is None:
        operators_map = OPERATORS
    params = build_base_params(requested_date=date_str, operators_map=operators_map) + [("page",1),("limit",1000)]
    r = requests.get(LIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT)
    r.raise_for_status()
    out = {}
    for item in r.json().get("items", []):
        oid = str(item.get("id") or "")
        if operators_map and oid not in operators_map:
            continue
        def to_int(value):
            try:
                return int(value or 0)
            except (TypeError, ValueError):
                return 0
        active = to_int(item.get("active"))
        dnd = to_int(item.get("dnd"))
        call = to_int(item.get("call"))
        ringing = to_int(item.get("ringing"))
        out[oid] = active + dnd + call + ringing
    return out

def extract_operators_from_calls(calls):
    operators = {}
    for call in calls:
        op = call.get("operator") or {}
        oid = str(op.get("id") or "")
        name = op.get("full_name") or op.get("fullName") or ""
        if oid and name:
            operators[oid] = name
    return operators


def fetch_new_numbers_total_by_noactive():
    if not USE_ACTIVE_CAMPAIGNS:
        return 0
    # Получаем исходный код функции fetch_new_numbers_total_by_active
    source = inspect.getsource(fetch_new_numbers_total_by_active)
    
    # Ищем все campaign_ids[] в формате ("campaign_ids[]", "<число>")
    active_ids = set(map(int, re.findall(r'\("campaign_ids\[\]",\s*"(\d+)"\)', source)))

    # Генерируем ID от 0 до 100, кроме найденных
    all_ids = set(range(0, 101))
    noactive_ids = sorted(all_ids - active_ids)

    # Собираем параметры запроса
    params = [("statuses[]", "1")]
    for campaign_id in noactive_ids:
        params.append(("campaign_ids[]", str(campaign_id)))
    params.append(("page", 1))
    params.append(("limit", 1))

    r = requests.get(CONTACT_LIST_URL, params=params, headers=HEADERS, timeout=REQUEST_TIMEOUT)
    r.raise_for_status()
    data = r.json()
    return data.get("totalCount", len(data.get("items", [])))

sched = None
SYNC_LOCK = threading.Lock()
SYNC_IN_FLIGHT = set()
LAST_SYNC_TS = {}
SYNC_MIN_INTERVAL = 50
DB_READY = False
DB_READY_LOCK = threading.Lock()

def ensure_db_ready():
    global DB_READY
    if DB_READY:
        return
    with DB_READY_LOCK:
        if DB_READY:
            return
        init_db()
        DB_READY = True

@app.before_request
def ensure_app_ready():
    ensure_db_ready()

def trigger_background_sync(date_str):
    now = time.time()
    with SYNC_LOCK:
        last = LAST_SYNC_TS.get(date_str, 0)
        if date_str in SYNC_IN_FLIGHT:
            return False
        if now - last < SYNC_MIN_INTERVAL:
            return False
        SYNC_IN_FLIGHT.add(date_str)
        LAST_SYNC_TS[date_str] = now
    def runner():
        try:
            sync_day(date_str)
        finally:
            with SYNC_LOCK:
                SYNC_IN_FLIGHT.discard(date_str)
    threading.Thread(target=runner, daemon=True).start()
    return True

def sync_yesterday():
    tz = pytz.timezone("Europe/Samara")
    day = datetime.now(tz).date() - timedelta(days=1)
    try:
        result = sync_day(format_date(day))
        print(f"Daily sync completed: {result}")
    except Exception as e:
        print(f"Daily sync failed: {str(e)}")

def sync_existing_dates():
    days_setting = get_setting("nightly_sync_days")
    try:
        nightly_days = int(days_setting) if days_setting is not None else NIGHTLY_SYNC_DAYS
    except ValueError:
        nightly_days = NIGHTLY_SYNC_DAYS
    dates = list_saved_dates()
    if not dates:
        print("Nightly sync: no saved dates")
        return
    dates_sorted = sorted(dates, key=to_db_key)
    if nightly_days > 0:
        dates_sorted = dates_sorted[-nightly_days:]
    print(f"Nightly sync: {len(dates_sorted)} dates")
    for ds in dates_sorted:
        try:
            sync_day(ds)
        except Exception as e:
            print(f"Nightly sync failed for {ds}: {e}")

def init_scheduler():
    global sched
    if sched is None:
        sched = BackgroundScheduler(timezone="Europe/Samara")
        # Обновление списка активных проектов в 10:00 (если включено)
        if USE_ACTIVE_CAMPAIGNS:
            sched.add_job(fetch_active_campaigns, 'cron', hour=10, minute=0)
        # Автосинхронизация вчерашнего дня в 00:10
        sched.add_job(sync_yesterday, 'cron', hour=0, minute=10)
        # Полный синк ЦК из Google Sheets раз в день в 00:30
        sched.add_job(sync_ck_sheet_all, 'cron', hour=0, minute=30)
        # Ночной полный пересинк всех дат из БД
        try:
            hour, minute = [int(x) for x in NIGHTLY_SYNC_TIME.split(":", 1)]
        except Exception:
            hour, minute = 2, 30
        sched.add_job(sync_existing_dates, 'cron', hour=hour, minute=minute)
        sched.start()

def render_dashboard():
    today = TEST_DATE or datetime.now(pytz.timezone("Europe/Samara")).strftime("%d.%m.%Y")
    return render_template('index.html', today=today)

@app.route('/')
def index():
    return redirect('/dashboard')

@app.route('/dashboard')
def dashboard():
    return render_dashboard()

@app.route('/reports')
def reports():
    return render_template('reports.html')

@app.route('/report/data')
def report_data():
    start = request.args.get("start")
    end = request.args.get("end")
    branch = (request.args.get("branch") or "").lower()
    sync_flag = (request.args.get("sync") or "").strip().lower() in {"1", "true", "yes", "y"}
    missing_dates = get_missing_dates_range(start, end) if start and end else []
    if (REPORT_AUTO_SYNC or sync_flag) and start and end:
        ensure_range_synced(start, end)
    if branch == "moscow":
        if not start or not end:
            uly = get_report_data(start, end)
            if uly.get("range"):
                start = start or uly["range"]["start"]
                end = end or uly["range"]["end"]
        data = get_moscow_report_data_db(start, end)
        data["missing_dates"] = missing_dates
        return jsonify(data)

    data = get_report_data(start, end)
    rows = filter_report_rows(data.get("rows", []), "ulyanovsk")
    data["rows"] = rows
    data["totals"] = recalc_report_totals(rows)
    if branch == "ulyanovsk":
        data["missing_dates"] = missing_dates
        return jsonify(data)

    range_start = start
    range_end = end
    if data.get("range"):
        range_start = range_start or data["range"]["start"]
        range_end = range_end or data["range"]["end"]
    moscow = get_moscow_report_data_db(range_start, range_end)
    return jsonify({
        "range": data.get("range") or moscow.get("range"),
        "ulyanovsk": data,
        "moscow": moscow,
        "missing_dates": missing_dates
    })

@app.route('/report/export')
def report_export():
    fmt = (request.args.get("format") or "xlsx").lower()
    if fmt != "xlsx":
        return jsonify({"error": "format not supported"}), 400
    start = request.args.get("start")
    end = request.args.get("end")
    branch = (request.args.get("branch") or "").lower()
    sync_flag = (request.args.get("sync") or "").strip().lower() in {"1", "true", "yes", "y"}
    if (REPORT_AUTO_SYNC or sync_flag) and start and end:
        ensure_range_synced(start, end)
    data = get_report_data(start, end)
    uly_rows = filter_report_rows(data.get("rows", []), "ulyanovsk")
    uly_totals = recalc_report_totals(uly_rows)
    data["rows"] = uly_rows
    data["totals"] = uly_totals
    range_start = start
    range_end = end
    if data.get("range"):
        range_start = range_start or data["range"]["start"]
        range_end = range_end or data["range"]["end"]

    def write_uly_sheet(ws, rows, totals):
        ws.append([
            "Оператор", "Всего", "На линии", "Диалогов",
            "Инвесторы: Перевод", "Инвесторы: Согласие", "Инвесторы: ЦК",
            "Агенты: Перевод", "Агенты: Согласие", "Агенты: ЦА",
            "Общая", "Среднее, сек"
        ])
        for row in rows:
            total_sum = (
                row.get("cs20", 0)
                + row.get("cs8", 0)
                + row.get("inv_ck", 0)
                + row.get("ag_transfer", 0)
                + row.get("ag_agreement", 0)
                + row.get("ag_ca", 0)
            )
            ws.append([
                row["operator_name"],
                row["all"],
                format_hms(row.get("line", 0)),
                row["total"],
                row["cs20"],
                row["cs8"],
                row.get("inv_ck", 0),
                row.get("ag_transfer", 0),
                row.get("ag_agreement", 0),
                row.get("ag_ca", 0),
                total_sum,
                row["avg"]
            ])
        if totals:
            ws.append([])
            total_sum = (
                totals.get("cs20", 0)
                + totals.get("cs8", 0)
                + totals.get("inv_ck", 0)
                + totals.get("ag_transfer", 0)
                + totals.get("ag_agreement", 0)
                + totals.get("ag_ca", 0)
            )
            ws.append([
                "ИТОГО",
                totals.get("all", 0),
                format_hms(totals.get("line", 0)),
                totals.get("total", 0),
                totals.get("cs20", 0),
                totals.get("cs8", 0),
                totals.get("inv_ck", 0),
                totals.get("ag_transfer", 0),
                totals.get("ag_agreement", 0),
                totals.get("ag_ca", 0),
                total_sum,
                totals.get("avg", 0)
            ])

    def write_moscow_sheet(ws, rows, totals):
        ws.append(["Оператор", "Звонков от минуты", "Согласия", "Встреч проведено", "Успешные сделки", "Выручка"])
        for row in rows:
            ws.append([
                row["operator_name"],
                row.get("calls_1m", 0),
                row.get("agreements", 0),
                row.get("meetings", 0),
                row.get("deals", 0),
                row.get("revenue", 0)
            ])
        if totals:
            ws.append([])
            ws.append([
                "ИТОГО",
                totals.get("calls_1m", 0),
                totals.get("agreements", 0),
                totals.get("meetings", 0),
                totals.get("deals", 0),
                totals.get("revenue", 0)
            ])

    wb = Workbook()
    ws = wb.active
    period = data.get("range")
    filename = "report.xlsx"

    if branch == "moscow":
        ws.title = "Moscow"
        moscow = get_moscow_report_data_db(range_start, range_end)
        write_moscow_sheet(ws, moscow.get("rows", []), moscow.get("totals", {}))
        period = moscow.get("range")
    elif branch == "ulyanovsk":
        ws.title = "Ulyanovsk"
        write_uly_sheet(ws, data.get("rows", []), data.get("totals", {}))
    else:
        ws.title = "Ulyanovsk"
        write_uly_sheet(ws, data.get("rows", []), data.get("totals", {}))
        moscow = get_moscow_report_data_db(range_start, range_end)
        ws_msk = wb.create_sheet("Moscow")
        write_moscow_sheet(ws_msk, moscow.get("rows", []), moscow.get("totals", {}))
        if not period:
            period = moscow.get("range")

    if period:
        filename = f"report_{period['start']}_{period['end']}.xlsx"
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route('/admin')
def admin_page():
    if not require_admin():
        return jsonify({"error": "unauthorized"}), 401
    return send_file(os.path.join(app.root_path, 'templates', 'admin.html'))

def require_admin():
    token = request.args.get("token") or ""
    # Reload from env to avoid stale token after .env edits
    env_token = os.getenv("CALLCENTER_ADMIN_TOKEN") or ""
    if not env_token:
        load_env()
        env_token = os.getenv("CALLCENTER_ADMIN_TOKEN") or ""
    if not env_token or token != env_token:
        return False
    return True

@app.route('/admin/db')
def admin_db():
    if not require_admin():
        return jsonify({"error": "unauthorized"}), 401
    start = request.args.get("start")
    end = request.args.get("end")
    available = list_saved_dates()
    empty = list_empty_dates()
    missing = []
    if start and end:
        start_date = parse_date(start)
        end_date = parse_date(end)
        if start_date > end_date:
            start_date, end_date = end_date, start_date
        available_set = set(available)
        empty_set = set(empty)
        d = start_date
        while d <= end_date:
            ds = format_date(d)
            if ds not in available_set and ds not in empty_set:
                missing.append(ds)
            d += timedelta(days=1)
        empty = [d for d in empty if parse_date(d) >= start_date and parse_date(d) <= end_date]
    return jsonify({
        "available": available,
        "missing": missing,
        "empty": empty
    })

@app.route('/admin/nightly')
def admin_nightly():
    if not require_admin():
        return jsonify({"error": "unauthorized"}), 401
    days = request.args.get("days")
    if days is not None:
        try:
            days_int = int(days)
        except ValueError:
            return jsonify({"error": "days must be int"}), 400
        set_setting("nightly_sync_days", days_int)
        return jsonify({"status": "ok", "days": days_int})
    current = get_setting("nightly_sync_days", str(NIGHTLY_SYNC_DAYS))
    try:
        current = int(current)
    except ValueError:
        current = NIGHTLY_SYNC_DAYS
    return jsonify({"days": current})

@app.route('/admin/sync')
def admin_sync():
    if not require_admin():
        return jsonify({"error": "unauthorized"}), 401
    date = request.args.get("date")
    start = request.args.get("start")
    end = request.args.get("end")
    try:
        if date:
            result = sync_day(date)
            return jsonify({"status": "ok", "result": result})
        if start and end:
            result = sync_range(start, end)
            return jsonify({"status": "ok", "result": result})
        return jsonify({"error": "date or start/end required"}), 400
    except Exception as e:
        import traceback
        error_details = traceback.format_exc()
        print(f"Admin sync failed: {str(e)}")
        print(error_details)
        return jsonify({"error": str(e), "details": error_details}), 500

@app.route('/admin/ck/sync')
def admin_ck_sync():
    if not require_admin():
        return jsonify({"error": "unauthorized"}), 401
    try:
        result = sync_ck_sheet_all()
        return jsonify({"status": "ok", "result": result})
    except Exception as e:
        import traceback
        error_details = traceback.format_exc()
        print(f"Admin CK sync failed: {str(e)}")
        print(error_details)
        return jsonify({"error": str(e), "details": error_details}), 500

@app.route('/stats')
def stats():
    def empty_stats_payload(server_error=None):
        payload = {
            "operators": {},
            "amo_operators": {},
            "status": {},
            "all": {},
            "line": {},
            "total": {},
            "cs8": {},
            "cs20": {},
            "ag_transfer": {},
            "ag_agreement": {},
            "inv_ck": {},
            "ag_ca": {},
            "lead_agent": {},
            "ck_lead": {},
            "avg": {},
            "amo_calls_1m": {},
            "amo_agreements": {},
            "amo_meetings": {},
            "amo_deals": {},
            "amo_revenue": {},
            "new": 0,
            "new_noactive": 0,
            "server_time": int(time.time() * 1000)
        }
        if server_error:
            payload["error"] = server_error
        return payload

    requested_date = request.args.get("date")
    print(f"Stats request args={dict(request.args)}")
    print(f"Stats request date={requested_date or 'today'}")
    if requested_date:
        today = datetime.now(pytz.timezone("Europe/Samara")).strftime("%d-%m-%Y")
        if requested_date == today:
            trigger_background_sync(requested_date)
            cached = get_day_stats_from_db(requested_date)
            if cached:
                return jsonify(cached)
            try:
                sync_day(requested_date)
            except Exception as e:
                print(f"Immediate sync failed for {requested_date}: {e}")
                return jsonify(empty_stats_payload(server_error=str(e)))
            cached = get_day_stats_from_db(requested_date)
            if cached:
                return jsonify(cached)
            return jsonify(empty_stats_payload(server_error="Нет данных после синхронизации"))

        # Прошедшие даты не пересинхронизируем на каждый запрос дашборда:
        # возвращаем срез из БД, а ручной пересинк остаётся через админку.
        cached = get_day_stats_from_db(requested_date)
        if cached:
            return jsonify(cached)

        try:
            sync_day(requested_date)
        except Exception as e:
            print(f"Stats sync failed for {requested_date}: {e}")
            return jsonify(empty_stats_payload(server_error=str(e)))
        cached = get_day_stats_from_db(requested_date)
        if cached:
            return jsonify(cached)
        return jsonify(empty_stats_payload(server_error=f"Нет данных за {requested_date}"))
    try:
        fetch_operators()
        calls   = fetch_all_calls_details(requested_date=requested_date, operators_map=OPERATORS)
        operators_from_calls = extract_operators_from_calls(calls)
        operators_map = OPERATORS or operators_from_calls
        total_dialogs = defaultdict(int)
        cs8     = fetch_counts(CS8, requested_date=requested_date, operators_map=operators_map)
        cs20    = fetch_counts(CS20, requested_date=requested_date, operators_map=operators_map)
        ag_transfer = fetch_counts(AG_TRANSFER, requested_date=requested_date, operators_map=operators_map)
        ag_agreement = fetch_counts(AG_AGREEMENT, requested_date=requested_date, operators_map=operators_map)
        lead_agent = fetch_counts(LEAD_AGENT, requested_date=requested_date, operators_map=operators_map)
        ck_counts = ck_counts_from_sheet(requested_date or datetime.now(pytz.timezone("Europe/Samara")).strftime("%d-%m-%Y"), operators_map) or Counter()
        allc    = fetch_all_counts(requested_date=requested_date, operators_map=operators_map)
        new_tot = fetch_new_numbers_total_by_active()
        new_noactive_tot = fetch_new_numbers_total_by_noactive()
        sums, cnts = defaultdict(int), defaultdict(int)
        for c in calls:
            oid = str(c.get("operator",{}).get("id") or "")
            if not operators_map or oid in operators_map:
                status_id = get_status_id(c)
                td = get_talk_duration(c)
                if status_id in STAT_FULL and td >= 20:
                    total_dialogs[oid] += 1
                    sums[oid] += td
                    cnts[oid] += 1
        avg    = {oid:(sums[oid]//cnts[oid] if cnts[oid] else 0) for oid in operators_map}
        status = fetch_current_status(requested_date=requested_date, operators_map=operators_map)
        line_times = fetch_line_times(requested_date=requested_date, operators_map=operators_map)

        active_operator_ids = {oid for oid, count in allc.items() if count > 0}
        operators_filtered = {oid: short_name(name) for oid, name in operators_map.items() if oid in active_operator_ids}

        ck_metric = {oid: ck_counts.get(oid, 0) for oid in operators_filtered}
        payload = {
            "operators": operators_filtered,
            "status":    {oid: status.get(oid) for oid in operators_filtered},
            "all":       {oid: allc.get(oid, 0) for oid in operators_filtered},
            "line":      {oid: line_times.get(oid, "00:00:00") for oid in operators_filtered},
            "total":     {oid: total_dialogs.get(oid, 0) for oid in operators_filtered},
            "cs8":       {oid: cs8.get(oid, 0) for oid in operators_filtered},
            "cs20":      {oid: cs20.get(oid, 0) for oid in operators_filtered},
            "ag_transfer": {oid: ag_transfer.get(oid, 0) for oid in operators_filtered},
            "ag_agreement": {oid: ag_agreement.get(oid, 0) for oid in operators_filtered},
            "inv_ck": dict(ck_metric),
            "ag_ca": {oid: 0 for oid in operators_filtered},
            "lead_agent": {oid: lead_agent.get(oid, 0) for oid in operators_filtered},
            "ck_lead":   ck_metric,
            "avg":       {oid: avg.get(oid, 0) for oid in operators_filtered},
            "new":       new_tot,
            "new_noactive": new_noactive_tot,
            "server_time": int(time.time() * 1000)
        }
        date_for_amo = requested_date or datetime.now(pytz.timezone(AMO_TZ)).strftime("%d-%m-%Y")
        return jsonify(merge_amo_counts(payload, date_for_amo))
    except Exception as e:
        print(f"Stats fallback failed for {requested_date or 'today'}: {e}")
        return jsonify(empty_stats_payload(server_error=str(e)))

if __name__ == '__main__':
    init_db()
    init_scheduler()
    app.run(host='0.0.0.0', port=8000)
